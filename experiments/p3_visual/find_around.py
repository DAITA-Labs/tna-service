"""Prototype: find_around_cell + find_around_range tools.

Generic "neighbourhood" primitives. Returns rich CellInfo per neighbour:
value (merge-resolved), dtype, merge anchor, style hints, etc.

Tests on real cells from MOPD/CB/63261 xlsx files to demonstrate.
"""
from __future__ import annotations

from dataclasses import dataclass, field, asdict
from datetime import datetime, date
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter


# =============================================================================
# Data models
# =============================================================================

@dataclass
class CellInfo:
    """Rich info for one cell."""
    coord:          str
    row:            int
    col:            str                  # column letter
    value:          Any                  # merge-resolved value
    raw_value:      Any                  # value at this exact cell (None if merged continuation)
    dtype:          str                  # "date" | "int" | "float" | "str" | "blank" | "formula"
    is_merged:      bool                 # True if this cell is inside a merge range
    is_merge_anchor:bool                 # True if this is the top-left of a merge
    merge_anchor:   Optional[str]        # the top-left coord (e.g., "B2")
    merge_range:    Optional[tuple]      # (min_row, min_col, max_row, max_col) if merged
    merge_size:     Optional[tuple]      # (rows, cols) span of the merge if merged
    is_bold:        bool
    fill_color:     Optional[str]        # hex like "FFCCCC"
    has_border:     bool


@dataclass
class CellNeighbourhood:
    """The 8 neighbours of a centre cell + the centre itself."""
    center:       CellInfo
    above:        Optional[CellInfo]
    below:        Optional[CellInfo]
    left:         Optional[CellInfo]
    right:        Optional[CellInfo]
    above_left:   Optional[CellInfo]
    above_right:  Optional[CellInfo]
    below_left:   Optional[CellInfo]
    below_right:  Optional[CellInfo]

    def as_grid_str(self) -> str:
        """Render the 3x3 neighbourhood as a text grid."""
        def fmt(c):
            if c is None: return "(out)"
            if c.value in (None, ""): return f"({c.coord} blank)"
            v = str(c.value)
            if len(v) > 22: v = v[:21] + "…"
            tag = ""
            if c.is_merge_anchor: tag += "🔝"
            elif c.is_merged: tag += "🟪"
            return f"{c.coord}={v}{tag}"
        rows = [
            [self.above_left, self.above, self.above_right],
            [self.left, self.center, self.right],
            [self.below_left, self.below, self.below_right],
        ]
        return "\n".join("  ".join(f"{fmt(c):<30}" for c in row) for row in rows)


@dataclass
class RangeNeighbourhood:
    """Cells around a range (the frame outside it)."""
    range_coord:   str                              # "C5:F8"
    above_row:     list[CellInfo]                   # cells directly above the range, same cols
    below_row:     list[CellInfo]
    left_col:      list[CellInfo]                   # cells directly left, same rows
    right_col:     list[CellInfo]
    corners:       dict[str, Optional[CellInfo]]    # NW, NE, SW, SE
    inside_summary: dict[str, Any]                  # dtype distribution + density of cells inside


# =============================================================================
# Build merge-resolved grid with rich cell info
# =============================================================================

def _detect_dtype(v: Any) -> str:
    if v is None or v == "": return "blank"
    if isinstance(v, (datetime, date)): return "date"
    if isinstance(v, bool): return "bool"
    if isinstance(v, int): return "int"
    if isinstance(v, float): return "float"
    if isinstance(v, str):
        s = v.strip()
        if s.startswith("="): return "formula"
        return "str"
    return "other"


def build_rich_grid(sh):
    """Return {(row, col_letter): CellInfo}, with merges resolved."""
    grid = {}
    # First pass: read direct cell properties for non-merged cells
    for row in sh.iter_rows(values_only=False):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue   # handled below via merge_cells.ranges
            raw_value = cell.value
            fill_color = None
            if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb:
                rgb = cell.fill.fgColor.rgb
                if isinstance(rgb, str) and len(rgb) == 8 and rgb != "00000000":
                    fill_color = rgb[2:]   # strip alpha
            is_bold = bool(cell.font and cell.font.bold)
            has_border = bool(cell.border and any(
                bool(getattr(cell.border, s) and getattr(cell.border, s).border_style)
                for s in ("left", "right", "top", "bottom")
            ))
            grid[(cell.row, cell.column_letter)] = CellInfo(
                coord=cell.coordinate,
                row=cell.row,
                col=cell.column_letter,
                value=raw_value,
                raw_value=raw_value,
                dtype=_detect_dtype(raw_value),
                is_merged=False,
                is_merge_anchor=False,
                merge_anchor=None,
                merge_range=None,
                merge_size=None,
                is_bold=is_bold,
                fill_color=fill_color,
                has_border=has_border,
            )
    # Second pass: handle merged ranges
    for mr in sh.merged_cells.ranges:
        top_coord = f"{get_column_letter(mr.min_col)}{mr.min_row}"
        top_value = sh.cell(mr.min_row, mr.min_col).value
        mr_tuple = (mr.min_row, mr.min_col, mr.max_row, mr.max_col)
        mr_size = (mr.max_row - mr.min_row + 1, mr.max_col - mr.min_col + 1)
        for r in range(mr.min_row, mr.max_row + 1):
            for c_i in range(mr.min_col, mr.max_col + 1):
                col_letter = get_column_letter(c_i)
                key = (r, col_letter)
                is_anchor = (r == mr.min_row and c_i == mr.min_col)
                # If anchor, update existing entry
                if is_anchor and key in grid:
                    grid[key].is_merged = True
                    grid[key].is_merge_anchor = True
                    grid[key].merge_anchor = top_coord
                    grid[key].merge_range = mr_tuple
                    grid[key].merge_size = mr_size
                else:
                    # Create a synthetic entry for the merge-continuation cell
                    grid[key] = CellInfo(
                        coord=f"{col_letter}{r}",
                        row=r,
                        col=col_letter,
                        value=top_value,                  # merge-resolved
                        raw_value=None,                   # this cell itself is empty
                        dtype=_detect_dtype(top_value),
                        is_merged=True,
                        is_merge_anchor=False,
                        merge_anchor=top_coord,
                        merge_range=mr_tuple,
                        merge_size=mr_size,
                        is_bold=False,
                        fill_color=None,
                        has_border=False,
                    )
    return grid


# =============================================================================
# Tools
# =============================================================================

def find_around_cell(grid: dict, cell_coord: str) -> CellNeighbourhood:
    """Return the 8 cells around `cell_coord` plus the cell itself."""
    # parse coord like "B4"
    col_letter, row = _split_coord(cell_coord)
    col_idx = column_index_from_string(col_letter)
    center = grid.get((row, col_letter))
    if center is None:
        # Construct an empty centre
        center = CellInfo(coord=cell_coord, row=row, col=col_letter, value=None,
                          raw_value=None, dtype="blank", is_merged=False,
                          is_merge_anchor=False, merge_anchor=None,
                          merge_range=None, merge_size=None,
                          is_bold=False, fill_color=None, has_border=False)
    def at(dr, dc):
        r2 = row + dr
        c2 = get_column_letter(col_idx + dc) if (col_idx + dc) > 0 else None
        if c2 is None or r2 < 1: return None
        return grid.get((r2, c2))
    return CellNeighbourhood(
        center=center,
        above=at(-1, 0),
        below=at(1, 0),
        left=at(0, -1),
        right=at(0, 1),
        above_left=at(-1, -1),
        above_right=at(-1, 1),
        below_left=at(1, -1),
        below_right=at(1, 1),
    )


def find_around_range(grid: dict, row_range: tuple[int, int], col_range: tuple[str, str]) -> RangeNeighbourhood:
    """Return the frame around a rectangular range + summary of cells inside."""
    r0, r1 = row_range
    c0_letter, c1_letter = col_range
    c0_idx = column_index_from_string(c0_letter)
    c1_idx = column_index_from_string(c1_letter)

    def at(r, c_i):
        if r < 1 or c_i < 1: return None
        return grid.get((r, get_column_letter(c_i)))

    # cells directly above the range (row r0-1, cols c0..c1)
    above_row = [at(r0 - 1, ci) for ci in range(c0_idx, c1_idx + 1)] if r0 > 1 else []
    # cells directly below
    below_row = [at(r1 + 1, ci) for ci in range(c0_idx, c1_idx + 1)]
    # cells directly left
    left_col = [at(r, c0_idx - 1) for r in range(r0, r1 + 1)] if c0_idx > 1 else []
    # cells directly right
    right_col = [at(r, c1_idx + 1) for r in range(r0, r1 + 1)]
    # corners
    corners = {
        "NW": at(r0 - 1, c0_idx - 1) if (r0 > 1 and c0_idx > 1) else None,
        "NE": at(r0 - 1, c1_idx + 1) if r0 > 1 else None,
        "SW": at(r1 + 1, c0_idx - 1) if c0_idx > 1 else None,
        "SE": at(r1 + 1, c1_idx + 1),
    }
    # inside summary: dtype distribution + density
    inside_cells = [at(r, ci) for r in range(r0, r1 + 1) for ci in range(c0_idx, c1_idx + 1)]
    inside_cells = [c for c in inside_cells if c is not None]
    total = len(inside_cells)
    non_blank = sum(1 for c in inside_cells if c.value not in (None, ""))
    dtype_counts = {}
    for c in inside_cells:
        dtype_counts[c.dtype] = dtype_counts.get(c.dtype, 0) + 1
    inside_summary = {
        "total_cells": total,
        "non_blank":   non_blank,
        "density":     round(non_blank / total, 3) if total else 0.0,
        "dtype_distribution": dtype_counts,
    }
    return RangeNeighbourhood(
        range_coord=f"{c0_letter}{r0}:{c1_letter}{r1}",
        above_row=above_row,
        below_row=below_row,
        left_col=left_col,
        right_col=right_col,
        corners=corners,
        inside_summary=inside_summary,
    )


def _split_coord(coord: str) -> tuple[str, int]:
    """Split 'AB123' → ('AB', 123)."""
    for i, ch in enumerate(coord):
        if ch.isdigit():
            return coord[:i], int(coord[i:])
    return coord, 1


# =============================================================================
# Demo on real files
# =============================================================================

def demo():
    print("=" * 75)
    print("DEMO 1 — DKN: cell V3 ('Planned' marker for PPS Submission stage)")
    print("=" * 75)
    wb = load_workbook("dataset/20260129 DKN AW26 DROP 2 WOMEN NOS CK PRO STATUS.xlsx",
                       data_only=True)
    sh = wb.active
    grid = build_rich_grid(sh)
    nb = find_around_cell(grid, "V3")
    print("\nNeighbourhood (3×3 visual):")
    print(nb.as_grid_str())
    print(f"\ncenter:      {nb.center.coord} value={nb.center.value!r} dtype={nb.center.dtype} bold={nb.center.is_bold}")
    print(f"above (V2):  value={nb.above.value!r} dtype={nb.above.dtype}   ← STAGE NAME captured by find_text_above")
    print(f"right (W3):  value={nb.right.value!r}    ← typical Actual sub-col")
    print(f"below (V4):  value={nb.below.value!r} dtype={nb.below.dtype}  ← first PLI's plan_date")

    print()
    print("=" * 75)
    print("DEMO 2 — CHRISTIAN BERG: cell B5 (vertical-merged identifier col)")
    print("=" * 75)
    wb = load_workbook("dataset/CHRISTIAN BERG- T&A.xlsx", data_only=True)
    sh = wb.active
    grid = build_rich_grid(sh)
    nb = find_around_cell(grid, "B5")
    print("\nNeighbourhood (3×3 visual):")
    print(nb.as_grid_str())
    print(f"\ncenter:           {nb.center.coord} value={nb.center.value!r}")
    print(f"  is_merged:       {nb.center.is_merged}")
    print(f"  is_merge_anchor: {nb.center.is_merge_anchor}")
    print(f"  merge_anchor:    {nb.center.merge_anchor}   ← top-left of merge")
    print(f"  merge_range:     {nb.center.merge_range}    ← (r0,c0,r1,c1)")
    print(f"  merge_size:      {nb.center.merge_size}     ← (rows, cols) span")

    print()
    print("=" * 75)
    print("DEMO 3 — 63261-TNA: cell B9 ('Plan' anchor for SHEET_IS_PLI strip)")
    print("=" * 75)
    wb = load_workbook("dataset/63261-TNA.xlsx", data_only=True)
    sh = wb.active
    grid = build_rich_grid(sh)
    nb = find_around_cell(grid, "B9")
    print("\nNeighbourhood (3×3 visual):")
    print(nb.as_grid_str())
    print(f"\ncenter (B9):  value={nb.center.value!r}        ← single Plan anchor (Case B)")
    print(f"right (C9):   value={nb.right.value!r} dtype={nb.right.dtype}  ← first horizontal date")
    print(f"above (B8):   value={nb.above.value!r}      ← blank; stage names at C8..")

    print()
    print("=" * 75)
    print("DEMO 4 — find_around_range on the CB FABRIC stage band (cols N-P, rows 1-2)")
    print("=" * 75)
    wb = load_workbook("dataset/CHRISTIAN BERG- T&A.xlsx", data_only=True)
    sh = wb.active
    grid = build_rich_grid(sh)
    rb = find_around_range(grid, row_range=(1, 2), col_range=("N", "P"))
    print(f"\nRange: {rb.range_coord}")
    print("Above row (cells above the band):")
    for c in rb.above_row:
        if c: print(f"  {c.coord}: value={c.value!r}")
    print("Below row (cells below the band — likely sub-headers PLAN/RECVD/APPD):")
    for c in rb.below_row:
        if c: print(f"  {c.coord}: value={c.value!r} dtype={c.dtype}")
    print(f"\nInside summary: {rb.inside_summary}")


if __name__ == "__main__":
    demo()
