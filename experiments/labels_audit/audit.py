"""Full labels audit driver — collects per-file findings and emits JSON corrections.

Run with:
  .venv/bin/python -m experiments.labels_audit.audit

Writes:
  experiments/labels_audit_corrections.json
"""
from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
DATASET = ROOT / "dataset"
EXTRACTED = DATASET / "extracted"
OUT = ROOT / "experiments" / "labels_audit_corrections.json"


# -----------------------------------------------------------------------------
# Six representative files (basename → xlsx path)
# -----------------------------------------------------------------------------

REP_FILES = [
    "20260129 DKN AW26 DROP 2 WOMEN NOS CK PRO STATUS",
    "CHRISTIAN BERG- T&A",
    "20260304 MOPD W26(1) MANOS COMPASS PRO",
    "63261-TNA",
    "new Eastman TnAs",
    "GUESS ATHLEISURE - MAIN FALL 26 MASTER FILE #1",
]


# -----------------------------------------------------------------------------
# helpers
# -----------------------------------------------------------------------------


def cellref(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def norm(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.isoformat()[:10]
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return str(v)
    return str(v).strip()


def merge_owner_map(ws) -> dict[tuple[int, int], tuple[int, int]]:
    owner: dict[tuple[int, int], tuple[int, int]] = {}
    for mr in ws.merged_cells.ranges:
        for row in range(mr.min_row, mr.max_row + 1):
            for col in range(mr.min_col, mr.max_col + 1):
                owner[(row, col)] = (mr.min_row, mr.min_col)
    return owner


def cell_value(ws, row: int, col: int, owner: dict | None = None):
    v = ws.cell(row=row, column=col).value
    if v is None and owner is not None and (row, col) in owner:
        o = owner[(row, col)]
        v = ws.cell(row=o[0], column=o[1]).value
    return v


def find_columns_by_header(ws, header_rows: Iterable[int]) -> dict[int, str]:
    """Return col_index -> joined header text across the listed rows."""
    out = {}
    for c in range(1, ws.max_column + 1):
        parts = []
        for r in header_rows:
            v = ws.cell(row=r, column=c).value
            if v is not None:
                parts.append(norm(v))
        if parts:
            out[c] = " | ".join(parts)
    return out


def value_in_row_match(ws, row: int, target: str, owner: dict | None = None) -> list[int]:
    """Return columns where the cell value matches the target string."""
    target_n = norm(target)
    cols = []
    for c in range(1, ws.max_column + 1):
        v = cell_value(ws, row, c, owner)
        if v is None:
            continue
        if norm(v) == target_n:
            cols.append(c)
    return cols


# -----------------------------------------------------------------------------
# Per-file analyses (driven by hand-coded knowledge of each file)
# -----------------------------------------------------------------------------


def analyze_dkn_w26_2026_01_29():
    """20260129 DKN AW26 DROP 2 WOMEN NOS CK PRO STATUS — ROW_PER_PLI."""
    name = "20260129 DKN AW26 DROP 2 WOMEN NOS CK PRO STATUS"
    wb = load_workbook(DATASET / f"{name}.xlsx", data_only=True)
    ws = wb["Sheet 1"]
    headers = find_columns_by_header(ws, [2, 3])
    labels = json.loads((EXTRACTED / f"{name}.json").read_text())

    findings = {
        "name": name,
        "sheet_count": 1,
        "data_sheets": ["Sheet 1"],
        "header_row_table": headers,
        "label_pli_count": labels["total_plis"],
        "data_rows": [4, 5, 6],
        "issues": [],
        "missing_fields": [],
        "stage_findings": [],
    }

    # Build per-PLI corrections
    corrections = []
    for i, pli in enumerate(labels["plis"]):
        row = 4 + i
        c = {"file": name, "pli_index": i, "source_row": row, "corrections": []}
        # io_number — label says e.g. "131673"; xlsx K (Buyer Po No) = 131673
        # there is no IO No column; F (Style No) has the supplier-internal code
        io_label = pli.get("io_number")
        po_no_val = norm(ws.cell(row=row, column=11).value)  # K
        if io_label and po_no_val == norm(io_label):
            c["corrections"].append({
                "canonical": "io_number",
                "old_value": io_label,
                "new_value": None,
                "reason": (
                    f"label value sourced from K{row} (header 'Buyer Po No') — that is "
                    "buyer_po_no, not io_number. No 'IO No' column exists in this sheet."
                ),
            })
            c["corrections"].append({
                "canonical": "buyer_po_no",
                "old_value": None,
                "new_value": io_label,
                "reason": "the value labeled as io_number is actually buyer's PO (K column)",
            })

        # fabric_code (label) — long composition string. Column I is "Fabric Quality".
        # Per spec: fabric_code should be a short code (max_len=20); long descriptive value belongs in fabric_name.
        fc = pli.get("fabric_code")
        if fc and len(fc) > 20:
            c["corrections"].append({
                "canonical": "fabric_code",
                "old_value": fc,
                "new_value": None,
                "reason": f"value at I{row} is a long fabric composition (>20 chars); spec says this is fabric_name",
            })
            c["corrections"].append({
                "canonical": "fabric_name",
                "old_value": None,
                "new_value": fc,
                "reason": "I header is 'Fabric Quality' — descriptive composition → fabric_name",
            })

        # color_code 6602 — column L header is "Color" (single column); spec allows code if dtype matches.
        # Label uses color_code=6602 (int_medium); OK to keep, but note that 'Color' header is ambiguous.

        # delivery_date — label says 2026-06-10; xlsx P column = "Etd Ex factory as per P.O" P4=10-JUN-2026
        # Spec anti-pattern says "Ex Factory" is metadata, not delivery_date.
        dd = pli.get("delivery_date")
        if dd:
            c["corrections"].append({
                "canonical": "delivery_date",
                "old_value": dd,
                "new_value": None,
                "reason": (
                    f"label value sourced from P{row} (header 'Etd Ex factory as per P.O'); "
                    "spec anti-pattern says 'Ex Factory' is metadata.ex_factory_date, not delivery_date"
                ),
            })
            c["corrections"].append({
                "canonical": "ex_factory_date",
                "old_value": None,
                "new_value": dd,
                "reason": "P column header 'Etd Ex factory' → metadata.ex_factory_date",
            })

        corrections.append(c)

    # missing fields
    findings["missing_fields"].append({
        "field": "ex_factory_date (metadata)",
        "evidence": "column O header 'Original order receipt Date' and column P 'Etd Ex factory as per P.O' both untracked",
    })

    # stage check — labels capture 6 stages (Trims Inhouse, Fabric Inhouse, PPS Submission, Sewing, Inspection, Ex Factory Shipment)
    # xlsx has stage bands: Trims Inhouse (R/S), Fabric Inhouse (T/U), PPS Submission (V/W/X), Cut Qty(Y), Sewing Start (Z/AA),
    # Sewing End (AB/AC), Sewing Qty (AD), Inspection band, Ex Factory band.
    # Labels collapse Sewing Start + Sewing End into single "Sewing" with planned=2026-05-23 = Z value (Sewing Start Plan).
    findings["stage_findings"].append({
        "issue": "stages compression",
        "detail": (
            "xlsx has separate Sewing Start (cols Z/AA) and Sewing End (AB/AC) bands; labels "
            "collapse into single 'Sewing' stage with planned=Sewing Start Plan only — loses end date"
        ),
    })

    return findings, corrections


def analyze_cb():
    name = "CHRISTIAN BERG- T&A"
    wb = load_workbook(DATASET / f"{name}.xlsx", data_only=True)
    ws = wb["CHRISTIAN BERG"]
    headers = find_columns_by_header(ws, [2, 3])
    labels = json.loads((EXTRACTED / f"{name}.json").read_text())
    owner = merge_owner_map(ws)

    findings = {
        "name": name,
        "sheet_count": 1,
        "data_sheets": ["CHRISTIAN BERG"],
        "header_row_table": headers,
        "label_pli_count": labels["total_plis"],
        "data_rows_per_pli": [
            (1, "1063 group", [4, 5, 6, 7]),
            (2, "1064 group", [9, 10, 11]),
        ],
        "issues": [],
        "missing_fields": [],
        "stage_findings": [],
    }

    corrections = []
    for i, pli in enumerate(labels["plis"]):
        src_rows = pli.get("source_rows", [])
        anchor_row = src_rows[-1] if src_rows else None
        c = {"file": name, "pli_index": i, "source_rows": src_rows, "corrections": []}

        # io_number — column B = "IO NO". Verify
        io_label = pli.get("io_number")
        if anchor_row:
            b_val = norm(cell_value(ws, anchor_row, 2, owner))
            if io_label and norm(io_label) == b_val:
                pass  # CORRECT
            else:
                c["corrections"].append({
                    "canonical": "io_number",
                    "old_value": io_label,
                    "new_value": b_val,
                    "reason": f"mismatch: B{anchor_row}={b_val!r} but label says {io_label!r}",
                })

        # style_code — label has long multiline "T-SLANIA LONG BP, ... D-T-SHIRT 3/4"
        # which is the F column raw value (multi-line cell concatenating style+description).
        # The proper code column is G ("ARTICLE") = 569510267 — labels skip it.
        sc = pli.get("style_code")
        if anchor_row:
            f_val = norm(cell_value(ws, anchor_row, 6, owner))  # F = STYLE
            g_val = norm(cell_value(ws, anchor_row, 7, owner))  # G = ARTICLE
            if sc and sc.replace(" ", "") == f_val.replace(" ", "").replace("\n", ""):
                c["corrections"].append({
                    "canonical": "style_code",
                    "old_value": sc,
                    "new_value": g_val,
                    "reason": (
                        f"label value sourced from F (header 'STYLE') — that cell is the combined "
                        f"style label, but the proper compact style_code is G{anchor_row}={g_val} "
                        "(header 'ARTICLE'). Spec max_len=20 expects a code, not 40+ char free text."
                    ),
                })

        # fabric_code — label has 100% ORG COT (...), which is column I = "FABRIC" composition (descriptive).
        # Should be fabric_name; xlsx has no separate fabric code column.
        fc = pli.get("fabric_code")
        if fc and len(fc) > 20:
            c["corrections"].append({
                "canonical": "fabric_code",
                "old_value": fc,
                "new_value": None,
                "reason": "I header 'FABRIC' holds composition; >20 chars → fabric_name per spec",
            })
            c["corrections"].append({
                "canonical": "fabric_name",
                "old_value": None,
                "new_value": fc,
                "reason": "I header 'FABRIC' composition → fabric_name",
            })

        # color_code — label is e.g. '422 - MAGENTA' which is descriptive text — should be color_name
        cc = pli.get("color_code")
        if cc and re.search(r"[A-Za-z]{3,}", cc):
            c["corrections"].append({
                "canonical": "color_code",
                "old_value": cc,
                "new_value": None,
                "reason": "K header 'COLOR' value is descriptive text (e.g. 'NAVY'); spec says short codes only → color_name",
            })
            c["corrections"].append({
                "canonical": "color_name",
                "old_value": None,
                "new_value": cc,
                "reason": "descriptive color text → color_name",
            })

        # delivery_date label = 2026-05-05 = D column "EX FAC DATE" — spec says this is metadata.ex_factory_date
        dd = pli.get("delivery_date")
        if dd:
            c["corrections"].append({
                "canonical": "delivery_date",
                "old_value": dd,
                "new_value": None,
                "reason": "label value sourced from D 'EX FAC DATE' — spec says this is metadata.ex_factory_date",
            })
            c["corrections"].append({
                "canonical": "ex_factory_date",
                "old_value": None,
                "new_value": dd,
                "reason": "D 'EX FAC DATE' → metadata.ex_factory_date",
            })

        corrections.append(c)

    findings["missing_fields"].append({
        "field": "plan_qty / cut_qty / sewing_qty (per-stage metadata.quantity)",
        "evidence": "L=ORDER QTY, M=PLAN QTY, Y=CUTTING QTY, AD=FEEDING QTY, AI=SEWING QTY — labels only capture L",
    })
    findings["stage_findings"].append({
        "issue": "stage count + sub-columns",
        "detail": (
            "labels capture 7 stages with metadata.received_date/approval/actual_date — matches the xlsx 7 bands "
            "(FABRIC PLAN/RECVD/APPD, SIZE SET PLAN/ACT, LOT CARD PLAN/ACT, CUTTING/FEEDING/SEWING START PLAN/START ACT/END PLAN/END ACT/QTY, FI PLAN/SUB). "
            "However planned_date for CUTTING/FEEDING/SEWING uses START PLAN; END PLAN dropped"
        ),
    })

    return findings, corrections


def analyze_mopd():
    name = "20260304 MOPD W26(1) MANOS COMPASS PRO"
    wb = load_workbook(DATASET / f"{name}.xlsx", data_only=True)
    ws = wb["Sheet 1"]
    headers = find_columns_by_header(ws, [2, 3])
    labels = json.loads((EXTRACTED / f"{name}.json").read_text())
    owner = merge_owner_map(ws)

    findings = {
        "name": name,
        "sheet_count": 1,
        "data_sheets": ["Sheet 1"],
        "header_row_table": headers,
        "label_pli_count": labels["total_plis"],
        "data_rows": [4, 5, 6, 7, 8, 9],
        "issues": [],
        "missing_fields": [],
        "stage_findings": [],
    }

    corrections = []
    for i, pli in enumerate(labels["plis"]):
        src_rows = pli.get("source_rows", [])
        anchor_row = src_rows[-1] if src_rows else None
        c = {"file": name, "pli_index": i, "source_rows": src_rows, "corrections": []}

        # io_number — column K = "Buyer Po No"
        io_label = pli.get("io_number")
        if anchor_row:
            k_val = norm(cell_value(ws, anchor_row, 11, owner))  # K
            if io_label and norm(io_label) == k_val:
                c["corrections"].append({
                    "canonical": "io_number",
                    "old_value": io_label,
                    "new_value": None,
                    "reason": f"label value sourced from K{anchor_row} (header 'Buyer Po No') — that is buyer_po_no, not io_number",
                })
                c["corrections"].append({
                    "canonical": "buyer_po_no",
                    "old_value": None,
                    "new_value": io_label,
                    "reason": "K column 'Buyer Po No' → buyer_po_no",
                })

        # style_code — label is "DWJE MANOS 08 1000000099 5000007827" (35 chars).
        # Column E header is "Style No" — and E4 = "DWJE MANOS 08 1000000052 5000006286" matches.
        # Label is correct in source but exceeds spec max_len=20.
        # Column F is "Style Name" with value like 3000000763 (10-digit numeric code) — but label uses
        # H column "Style Description" for style_name.
        # So actually the file has a mismatch: F is labeled "Style Name" in xlsx but holds a numeric code.
        # The 10-digit numeric in F (3000000763, 3000001132, etc.) is more likely the supplier's style_code.
        # The compound code in E is closer to a description / buyer-internal reference.

        # fabric_code — label = long composition string → should be fabric_name
        fc = pli.get("fabric_code")
        if fc and len(fc) > 20:
            c["corrections"].append({
                "canonical": "fabric_code",
                "old_value": fc,
                "new_value": None,
                "reason": "I 'Fabric Quality' value > 20 chars; spec → fabric_name",
            })
            c["corrections"].append({
                "canonical": "fabric_name",
                "old_value": None,
                "new_value": fc,
                "reason": "I 'Fabric Quality' composition → fabric_name",
            })

        # color_code — '4139–NAVY TEAL', '1183 - SILKY WHITE' etc — these are 'code - name' compound values
        # Strictly per spec they're closer to color_name (descriptive); some files split them. Mark as
        # potentially conflated.
        cc = pli.get("color_code")
        if cc and re.search(r"[A-Za-z]{3,}", cc):
            c["corrections"].append({
                "canonical": "color_code",
                "old_value": cc,
                "new_value": None,
                "reason": "L 'Color' value is 'code-name' compound or descriptive; per spec → color_name (or split)",
            })
            c["corrections"].append({
                "canonical": "color_name",
                "old_value": None,
                "new_value": cc,
                "reason": "L 'Color' compound → color_name",
            })

        # delivery_date — column P "Etd Ex factory as per P.O" — spec anti-pattern says metadata
        dd = pli.get("delivery_date")
        if dd:
            c["corrections"].append({
                "canonical": "delivery_date",
                "old_value": dd,
                "new_value": None,
                "reason": "P 'Etd Ex factory as per P.O' → spec anti-pattern → metadata.ex_factory_date",
            })
            c["corrections"].append({
                "canonical": "ex_factory_date",
                "old_value": None,
                "new_value": dd,
                "reason": "P 'Etd Ex factory' → metadata.ex_factory_date",
            })

        corrections.append(c)

    findings["missing_fields"].append({
        "field": "style_name disambiguation",
        "evidence": (
            "xlsx F header 'Style Name' holds numeric code (e.g. 3000000763) and H 'Style Description' "
            "holds the descriptive phrase. Labels pick H for style_name (correct in intent) but F is "
            "left uncaptured — could be a second style_code-like field"
        ),
    })
    findings["stage_findings"].append({
        "issue": "stages not in labels view",
        "detail": "PLIs in labels capture 6 stages (Trims, Fabric, Sizeset Submission, Sewing, Inspection, Ex Factory) — same shape as DKN",
    })

    return findings, corrections


def analyze_63261():
    name = "63261-TNA"
    wb = load_workbook(DATASET / f"{name}.xlsx", data_only=True)
    ws = wb["Sheet1"]
    labels = json.loads((EXTRACTED / f"{name}.json").read_text())

    findings = {
        "name": name,
        "sheet_count": len(wb.sheetnames),
        "all_sheets": wb.sheetnames,
        "data_sheets": ["Sheet1"],
        "label_pli_count": labels["total_plis"],
        "issues": [],
        "missing_fields": [],
        "stage_findings": [],
    }

    corrections = []
    for i, pli in enumerate(labels["plis"]):
        c = {"file": name, "pli_index": i, "corrections": []}

        # io_number = '63261' — A4='Job No' B4=63261 — CORRECT. But 'Job No' alias not in spec.
        io_label = pli.get("io_number")
        if io_label == "63261":
            # OK
            pass

        # quantity = 16200 — A5='Quantity :' B5=16200 — CORRECT
        # delivery_date = 2026-05-17 — D5='Delivery date' E5=2026-05-17 — CORRECT

        # style_code = "" — empty string in label. Should be None (or removed).
        sc = pli.get("style_code")
        if sc == "":
            c["corrections"].append({
                "canonical": "style_code",
                "old_value": "",
                "new_value": None,
                "reason": "empty-string instead of null; xlsx has no style_code column for this PLI",
            })
        color_code = pli.get("color_code")
        if color_code == "":
            c["corrections"].append({
                "canonical": "color_code",
                "old_value": "",
                "new_value": None,
                "reason": "empty-string instead of null",
            })
        fc = pli.get("fabric_code")
        if fc == "":
            c["corrections"].append({
                "canonical": "fabric_code",
                "old_value": "",
                "new_value": None,
                "reason": "empty-string instead of null",
            })
        corrections.append(c)

    findings["missing_fields"].extend([
        {"field": "ex_factory_date (metadata)", "evidence": "D4 'Ex-Fty date' E4=2026-05-07 not captured"},
        {"field": "order_receipt_date (metadata)", "evidence": "D3 'Order receipt' E3=2026-02-06 not captured"},
        {"field": "order_l_d (metadata)", "evidence": "G3 'Order L/D' H3=90 not captured"},
        {"field": "report_date (metadata)", "evidence": "A3 'Date :' B3=2026-04-21 not captured"},
    ])
    findings["stage_findings"].append({
        "issue": "stages",
        "detail": "labels capture 11 stages across Pre-Production, Fabric, Production bands — matches xlsx",
    })

    return findings, corrections


def analyze_eastman():
    name = "new Eastman TnAs"
    wb = load_workbook(DATASET / f"{name}.xlsx", data_only=True)
    labels = json.loads((EXTRACTED / f"{name}.json").read_text())

    findings = {
        "name": name,
        "sheet_count": len(wb.sheetnames),
        "all_sheets": wb.sheetnames,
        "label_pli_count": labels["total_plis"],
        "issues": [],
        "missing_fields": [],
        "stage_findings": [],
    }

    # Inspect each sheet
    label_sheets = {p.get("source_sheet"): i for i, p in enumerate(labels["plis"])}
    data_sheets = []
    non_data_sheets = []
    sheet_audit = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        a4 = norm(ws["A4"].value) if ws.max_row >= 4 else ""
        b4 = norm(ws["B4"].value) if ws.max_row >= 4 else ""
        is_index_sheet = (a4.upper() == "S.NO" or "S.NO" in (a4.upper(),)) and (b4.upper() == "JOB NO")
        if a4 == "Job No":
            data_sheets.append(sn)
            # Compare to label
            label_idx = label_sheets.get(sn)
            if label_idx is None:
                sheet_audit.append({
                    "sheet": sn,
                    "status": "missing_label",
                    "xlsx_io": b4,
                })
                continue
            pli = labels["plis"][label_idx]
            label_io = pli.get("io_number")
            label_qty = pli.get("quantity")
            xlsx_qty = ws["B5"].value
            xlsx_deliv = ws["E5"].value
            label_deliv = pli.get("delivery_date")
            mismatches = []
            if norm(label_io) != b4:
                mismatches.append(f"io_number label={label_io!r} xlsx_B4={b4!r}")
            if label_qty != xlsx_qty:
                mismatches.append(f"quantity label={label_qty!r} xlsx_B5={xlsx_qty!r}")
            if isinstance(xlsx_deliv, (date, datetime)):
                xd = xlsx_deliv.isoformat()[:10]
            else:
                xd = norm(xlsx_deliv)
            if norm(label_deliv) != xd:
                mismatches.append(f"delivery_date label={label_deliv!r} xlsx_E5={xd!r}")
            sheet_audit.append({
                "sheet": sn,
                "status": "ok" if not mismatches else "mismatch",
                "mismatches": mismatches,
            })
        else:
            non_data_sheets.append(sn)

    findings["data_sheets"] = data_sheets
    findings["non_data_sheets"] = non_data_sheets
    findings["sheet_audit_summary"] = {
        "data_sheet_count": len(data_sheets),
        "non_data_sheet_count": len(non_data_sheets),
        "labeled_sheets": len(label_sheets),
        "all_data_sheets_labeled": all(s.get("status") == "ok" for s in sheet_audit),
    }
    findings["sheet_audit"] = sheet_audit
    findings["missing_fields"].append({
        "field": "ex_factory_date (metadata)",
        "evidence": "D4 'Ex-Fty date' present on every data sheet but never captured as metadata",
    })

    corrections = []
    # Empty strings should be null
    for i, pli in enumerate(labels["plis"]):
        c = {"file": name, "pli_index": i, "sheet": pli.get("source_sheet"), "corrections": []}
        for k in ("style_code", "color_code", "fabric_code"):
            if pli.get(k) == "":
                c["corrections"].append({
                    "canonical": k,
                    "old_value": "",
                    "new_value": None,
                    "reason": "empty-string instead of null",
                })
        if c["corrections"]:
            corrections.append(c)

    return findings, corrections


def analyze_guess1():
    name = "GUESS ATHLEISURE - MAIN FALL 26 MASTER FILE #1"
    wb = load_workbook(DATASET / f"{name}.xlsx", data_only=True)
    ws = wb["MAIN FALL 26"]
    headers = find_columns_by_header(ws, [1])
    labels = json.loads((EXTRACTED / f"{name}.json").read_text())

    findings = {
        "name": name,
        "sheet_count": len(wb.sheetnames),
        "all_sheets": wb.sheetnames,
        "data_sheets": ["MAIN FALL 26"],
        "header_row_table": headers,
        "label_pli_count": labels["total_plis"],
        "issues": [],
        "missing_fields": [],
        "stage_findings": [],
    }

    # Spot-check 5 PLIs
    corrections = []
    sample_indices = [0, 1, 10, 100, 169]
    for i in sample_indices:
        if i >= len(labels["plis"]):
            continue
        pli = labels["plis"][i]
        src_rows = pli.get("source_rows", [])
        if not src_rows:
            continue
        row = src_rows[0]
        c = {"file": name, "pli_index": i, "source_row": row, "corrections": []}

        # E column = ION, K column = PO NO
        e_val = norm(ws.cell(row=row, column=5).value)  # E = ION
        k_val = norm(ws.cell(row=row, column=11).value)  # K = PO NO
        z_val = ws.cell(row=row, column=26).value  # Z = EX FAC
        aa_val = ws.cell(row=row, column=27).value  # AA = EX CON

        io_label = pli.get("io_number")
        if io_label and norm(io_label) == k_val:
            c["corrections"].append({
                "canonical": "io_number",
                "old_value": io_label,
                "new_value": e_val,
                "reason": (
                    f"label value at K{row} (header 'PO NO') is buyer_po_no, not io_number. "
                    f"The 'ION' column E{row}={e_val} is the actual io_number (Internal Order Number alias)."
                ),
            })
            c["corrections"].append({
                "canonical": "buyer_po_no",
                "old_value": None,
                "new_value": io_label,
                "reason": "K 'PO NO' → buyer_po_no",
            })

        # fabric_code — long descriptive string
        fc = pli.get("fabric_code")
        if fc and len(fc) > 20:
            c["corrections"].append({
                "canonical": "fabric_code",
                "old_value": fc,
                "new_value": None,
                "reason": "H 'FABRIC' value > 20 chars descriptive → fabric_name",
            })
            c["corrections"].append({
                "canonical": "fabric_name",
                "old_value": None,
                "new_value": fc,
                "reason": "H 'FABRIC' composition → fabric_name",
            })

        # delivery_date label vs xlsx
        dd = pli.get("delivery_date")
        # Z = EX FAC; AA = EX CON. Z dates are typically ~3-7 days before AA (the consignee).
        def _to_iso(v):
            if isinstance(v, (date, datetime)):
                return v.isoformat()[:10]
            return norm(v) if v is not None else None
        z_iso = _to_iso(z_val)
        if dd and z_iso == dd:
            c["corrections"].append({
                "canonical": "delivery_date",
                "old_value": dd,
                "new_value": None,
                "reason": (
                    f"label value at Z{row} ('EX FAC' = ex-factory date) — spec anti-pattern → "
                    "this is metadata.ex_factory_date. AA 'EX CON' (ex-consignee) is more like delivery_date."
                ),
            })
            c["corrections"].append({
                "canonical": "ex_factory_date",
                "old_value": None,
                "new_value": dd,
                "reason": "Z 'EX FAC' → metadata.ex_factory_date",
            })

        if c["corrections"]:
            corrections.append(c)

    findings["missing_fields"].extend([
        {"field": "country / destination (metadata)", "evidence": "L 'COUNTRY' column"},
        {"field": "price (metadata)", "evidence": "J 'PRICE' column"},
        {"field": "test_request (metadata)", "evidence": "AB column"},
        {"field": "mats_po (metadata)", "evidence": "AC 'MATS PO#' column"},
        {"field": "ex_consignee_date (metadata)", "evidence": "AA 'EX CON' column (date of consignee handover)"},
    ])
    findings["stage_findings"].append({
        "issue": "missing stages",
        "detail": (
            "xlsx has columns AE/AF/AG ('PROGRAM SUBMIT ON', 'FABRIC ETA PLAN', 'FABRIC IN-HOUSED ON') "
            "and labels capture exactly these 3. But many xlsx stages are missing (no Cutting/Sewing/FI bands "
            "in this sheet — appears truncated). Stage coverage is correct given the sheet contents."
        ),
    })

    return findings, corrections


# -----------------------------------------------------------------------------
# main
# -----------------------------------------------------------------------------


def main():
    all_findings = []
    all_corrections = []
    for fn in (
        analyze_dkn_w26_2026_01_29,
        analyze_cb,
        analyze_mopd,
        analyze_63261,
        analyze_eastman,
        analyze_guess1,
    ):
        try:
            f, c = fn()
            all_findings.append(f)
            all_corrections.extend(c)
            print(f"OK: {f['name']}")
        except Exception as e:  # noqa
            print(f"FAILED {fn.__name__}: {e}")
            raise

    OUT.write_text(json.dumps(all_corrections, indent=2, default=str))
    print(f"wrote {len(all_corrections)} correction records to {OUT}")
    return all_findings, all_corrections


if __name__ == "__main__":
    main()
