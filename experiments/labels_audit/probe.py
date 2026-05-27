"""Audit helper: dump xlsx headers + sample data rows side-by-side with labels.

Run as a module:
  .venv/bin/python -m experiments.labels_audit.probe <basename_fragment>

Examples:
  python -m experiments.labels_audit.probe DKN
  python -m experiments.labels_audit.probe CHRISTIAN
  python -m experiments.labels_audit.probe Eastman --list-sheets

Read-only on dataset/.
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
DATASET = ROOT / "dataset"
EXTRACTED = DATASET / "extracted"


def find_xlsx(fragment: str) -> Path:
    fragment_lower = fragment.lower()
    matches = [
        p for p in DATASET.glob("*.xlsx") if fragment_lower in p.name.lower()
    ]
    if not matches:
        raise SystemExit(f"no xlsx found matching {fragment!r}")
    if len(matches) > 1:
        names = "\n  ".join(p.name for p in matches)
        raise SystemExit(f"multiple xlsx match {fragment!r}:\n  {names}")
    return matches[0]


def cellref(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def pretty(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.isoformat()[:10]
    s = str(v).strip()
    s = s.replace("\n", " ")
    if len(s) > 60:
        s = s[:57] + "..."
    return s


def merge_owner_map(ws) -> dict[tuple[int, int], tuple[int, int]]:
    """For every cell inside a merged range, map it -> top-left of the merge.

    Used for io columns with vertical merges (CHRISTIAN BERG-style).
    """
    owner: dict[tuple[int, int], tuple[int, int]] = {}
    for mr in ws.merged_cells.ranges:
        for row in range(mr.min_row, mr.max_row + 1):
            for col in range(mr.min_col, mr.max_col + 1):
                owner[(row, col)] = (mr.min_row, mr.min_col)
    return owner


def dump_sheet(ws, max_rows: int = 12, max_cols: int = 40) -> None:
    print(f"\n--- Sheet: {ws.title} ({ws.max_row} rows × {ws.max_column} cols) ---")
    owner = merge_owner_map(ws)
    nrows = min(ws.max_row, max_rows)
    ncols = min(ws.max_column, max_cols)

    header = [f"{get_column_letter(c):>4}" for c in range(1, ncols + 1)]
    print(" row | " + " | ".join(header))
    print(" --- + " + "-+-".join(["----"] * ncols))
    for r in range(1, nrows + 1):
        row_cells = []
        for c in range(1, ncols + 1):
            v = ws.cell(row=r, column=c).value
            if v is None and (r, c) in owner:
                ow = owner[(r, c)]
                v = ws.cell(row=ow[0], column=ow[1]).value
                if v is not None:
                    row_cells.append(f"·{pretty(v)[:20]}")  # mark merged-from
                else:
                    row_cells.append("")
            else:
                row_cells.append(pretty(v)[:20])
        print(f"{r:>4} | " + " | ".join(f"{c[:20]:>20}" for c in row_cells))


def load_labels(name: str) -> dict | None:
    p = EXTRACTED / f"{name}.json"
    if not p.exists():
        return None
    return json.loads(p.read_text())


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("fragment")
    ap.add_argument("--list-sheets", action="store_true")
    ap.add_argument("--sheet", help="dump a single sheet by name")
    ap.add_argument("--max-rows", type=int, default=12)
    ap.add_argument("--max-cols", type=int, default=40)
    ap.add_argument("--all-sheets", action="store_true")
    ap.add_argument("--labels", action="store_true", help="print labels summary")
    ap.add_argument("--full-row", type=int, help="print full content of a row (no truncation)")
    args = ap.parse_args()

    p = find_xlsx(args.fragment)
    print(f"file: {p.name}")
    wb = load_workbook(p, data_only=True, read_only=False)

    print(f"sheets ({len(wb.sheetnames)}): {wb.sheetnames}")

    if args.list_sheets:
        return

    sheets = wb.sheetnames if args.all_sheets else [args.sheet or wb.sheetnames[0]]
    for sheetname in sheets:
        ws = wb[sheetname]
        dump_sheet(ws, max_rows=args.max_rows, max_cols=args.max_cols)
        if args.full_row is not None:
            print(f"\nFull content of row {args.full_row} on '{sheetname}':")
            for c in range(1, ws.max_column + 1):
                v = ws.cell(row=args.full_row, column=c).value
                if v is not None:
                    print(f"  {cellref(args.full_row, c)} = {v!r}")

    if args.labels:
        labels = load_labels(p.stem)
        if labels is None:
            print(f"\n(no labels file for {p.stem})")
        else:
            print("\nLabels summary:")
            print(f"  total_plis: {labels.get('total_plis')}")
            for i, pli in enumerate(labels.get("plis", [])):
                ids = {
                    k: pli.get(k)
                    for k in (
                        "io_number", "style_code", "style_name", "color_code",
                        "color_name", "fabric_code", "fabric_name",
                        "delivery_date", "quantity",
                    )
                    if pli.get(k) is not None
                }
                print(
                    f"  PLI {i}: sheet={pli.get('source_sheet')!r} "
                    f"rows={pli.get('source_rows')} | {ids}"
                )


if __name__ == "__main__":
    main()
