"""Structure tools — merged regions, occupancy counts."""
from __future__ import annotations
from openpyxl.utils import column_index_from_string
from app.models.workbook import WorkbookCtx, MergedRegion
from app.repositories.workbook_tools._registry import tool


@tool("get_merged_regions")
def get_merged_regions(ctx: WorkbookCtx, sheet: str) -> list[MergedRegion]:
    """Return every merged region in `sheet` with its anchor value."""
    ws = ctx.wb[sheet]
    out: list[MergedRegion] = []
    for mr in ws.merged_cells.ranges:
        anchor_value = ws.cell(row=mr.min_row, column=mr.min_col).value
        anchor = ws.cell(row=mr.min_row, column=mr.min_col).coordinate
        out.append(MergedRegion(
            cell_range=mr.coord,
            anchor=anchor,
            anchor_value=anchor_value,
        ))
    return out


@tool("count_non_empty_rows_in_column")
def count_non_empty_rows_in_column(
    ctx: WorkbookCtx, sheet: str, column: str,
    row_range: tuple[int, int] | None = None,
) -> int:
    """Count rows in `column` (inside row_range) with a non-empty value."""
    ws = ctx.wb[sheet]
    col_idx = column_index_from_string(column)
    r0, r1 = row_range or (1, ws.max_row or 1)
    return sum(
        1 for r in range(r0, r1 + 1)
        if ws.cell(row=r, column=col_idx).value is not None
    )
