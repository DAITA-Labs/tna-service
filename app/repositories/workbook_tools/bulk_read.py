"""Bulk-read tools — bounded windows over the sheet."""
from __future__ import annotations
from datetime import date, datetime

from openpyxl.utils import get_column_letter

from app.enums.cell_dtype import CellDtype
from app.models.workbook import Cell, CellGrid, WorkbookCtx
from app.repositories.workbook_tools._registry import tool


def _infer_dtype(v: object) -> CellDtype:
    """Map a raw openpyxl cell value to its CellDtype tag."""
    if v is None:
        return CellDtype.EMPTY
    if isinstance(v, bool):
        return CellDtype.BOOL
    if isinstance(v, int):
        return CellDtype.INT
    if isinstance(v, float):
        return CellDtype.FLOAT
    if isinstance(v, (date, datetime)):
        return CellDtype.DATE
    if isinstance(v, str) and v.startswith("#") and v.endswith("!"):
        return CellDtype.ERROR
    return CellDtype.STR


def _cell(ws, row: int, col: int) -> Cell:
    raw = ws.cell(row=row, column=col).value
    addr = f"{get_column_letter(col)}{row}"
    return Cell(row=row, col=col, address=addr, value=raw, dtype=_infer_dtype(raw))


@tool("peek_sheet")
def peek_sheet(ctx: WorkbookCtx, sheet: str, rows: int = 10, cols: int = 15) -> CellGrid:
    """Top-left peek; bounded window — useful for header + early-data probes."""
    ws = ctx.wb[sheet]
    max_r = min(rows, ws.max_row or rows)
    max_c = min(cols, ws.max_column or cols)
    cells: list[Cell] = []
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            cell = _cell(ws, r, c)
            if cell.value is not None:
                cells.append(cell)
    return CellGrid(
        sheet=sheet,
        cell_range=f"A1:{get_column_letter(max_c)}{max_r}",
        cells=cells,
    )


@tool("sample_rows")
def sample_rows(ctx: WorkbookCtx, sheet: str, row_indices: list[int]) -> list[list[Cell]]:
    """Read a small set of rows fully (all non-empty columns)."""
    ws = ctx.wb[sheet]
    max_c = ws.max_column or 0
    out: list[list[Cell]] = []
    for r in row_indices:
        row_cells = [_cell(ws, r, c) for c in range(1, max_c + 1)]
        out.append([c for c in row_cells if c.value is not None])
    return out


@tool("read_range")
def read_range(
    ctx: WorkbookCtx, sheet: str,
    row_range: tuple[int, int], col_range: tuple[int, int],
) -> CellGrid:
    """Read every non-empty cell in [row_range] × [col_range], inclusive."""
    ws = ctx.wb[sheet]
    r0, r1 = row_range
    c0, c1 = col_range
    cells: list[Cell] = []
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            cell = _cell(ws, r, c)
            if cell.value is not None:
                cells.append(cell)
    return CellGrid(
        sheet=sheet,
        cell_range=f"{get_column_letter(c0)}{r0}:{get_column_letter(c1)}{r1}",
        cells=cells,
    )
