"""Orchestrates deterministic detectors to produce a SheetPlan per sheet.

SheetRowPlanner is the top of the planner pipeline: it surveys the sheet,
decides PLI mode from structural signals, detects stage bands and KV anchors,
classifies rows, segments PLI blocks when applicable, and assembles the final
SheetPlan artifact returned to the Haystack pipeline.
"""
from __future__ import annotations

from typing import Any

from haystack import component
from openpyxl.utils import get_column_letter

from app.core.logs import get_logger
from app.enums.pli_mode import PliMode
from app.enums.row_role import RowRole
from app.enums.stage_scope import StageScope
from app.models.artifacts import (
    HeaderLabel,
    KVAnchor,
    PliBlock,
    SheetPlan,
    SheetSignals,
    StageBandSpec,
    RowSpec,
)
from app.services.planner.block_segmenter import segment_blocks
from app.services.planner.kv_anchor_detector import detect_kv_anchors
from app.services.planner.row_classifier import classify_rows
from app.services.planner.stage_band_detector import detect_stage_bands
from app.services.planner.surveyor import survey_sheet

log = get_logger(__name__)


def _decide_pli_mode(signals: SheetSignals) -> PliMode:
    """Decide PLI mode from structural sheet signals.

    Distinguishes ROW_PER_PLI, SHEET_IS_PLI, and SECTION_PER_PLI by examining
    identity column candidates against KV label hit positions and blank-run gaps.
    """
    # A column is a real tabular identity column (header) when:
    #  (a) all its KV hits fall on a single row (the header row), OR
    #  (b) all its KV hits carry the same label text (repeat-header pattern).
    # If hits span multiple distinct rows AND carry different labels, the column
    # is a scattered KV-label column — a signal for SHEET_IS_PLI.
    real_identity_cols: list[str] = []
    for col in signals.identity_col_candidates:
        col_kv_hits: list[tuple[str, int]] = [
            (label, int(c[len(col):]))
            for label, c in signals.kv_label_hits
            if c.startswith(col) and c[len(col):].isdigit()
        ]
        if not col_kv_hits:
            continue
        col_kv_rows = [row for _, row in col_kv_hits]
        col_kv_labels = [label for label, _ in col_kv_hits]
        # Single-row concentration OR repeated-header (same label at many rows) →
        # tabular identity column.  Different labels at different rows → KV scatter.
        if len(set(col_kv_rows)) == 1 or len(set(col_kv_labels)) == 1:
            real_identity_cols.append(col)

    has_identity_col = len(real_identity_cols) >= 1
    n_kv_hits = len(signals.kv_label_hits)
    n_blank_gaps = len(signals.blank_run_gaps)

    if has_identity_col:
        if n_blank_gaps >= 2 and n_kv_hits >= 3:
            return PliMode.SECTION_PER_PLI
        return PliMode.ROW_PER_PLI
    if n_kv_hits >= 2:
        return PliMode.SHEET_IS_PLI
    return PliMode.ROW_PER_PLI


def _pick_identity_column(signals: SheetSignals) -> str | None:
    """Return the best identity column letter from candidates, or None."""
    if not signals.identity_col_candidates:
        return None
    for col in signals.identity_col_candidates:
        for hit in signals.header_vocab_hits.get(col, []):
            h = hit.lower()
            if "io" in h or "job" in h or "buyer po" in h:
                return col
    return signals.identity_col_candidates[0]


def _pick_quantity_column(workbook_ctx: Any, sheet: str, signals: SheetSignals) -> str | None:
    """Scan the first ten rows for a quantity-header cell and return its column letter."""
    ws = workbook_ctx.wb[sheet]
    for c in range(1, signals.max_col + 1):
        for r in range(1, min(signals.max_row, 10) + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str):
                vn = " ".join(v.strip().lower().split())
                if any(t in vn for t in ("qty", "quantity", "order qty", "plan qty")):
                    return get_column_letter(c)
    return None


def _classify_sheet_rows(
    workbook_ctx: Any,
    sheet: str,
    signals: SheetSignals,
    pli_mode: PliMode,
    identity_column: str | None,
) -> list[RowSpec]:
    """Classify sheet rows into RowSpec entries, or return empty for SHEET_IS_PLI."""
    if pli_mode is PliMode.SHEET_IS_PLI:
        return []
    qty_col = _pick_quantity_column(workbook_ctx, sheet, signals)
    return classify_rows(
        workbook_ctx, sheet, signals,
        identity_column=identity_column,
        quantity_column_hint=qty_col,
    )


def _segment_pli_blocks_if_applicable(
    rows: list[RowSpec],
    kv_anchors: list[KVAnchor],
    signals: SheetSignals,
    stage_bands: list[StageBandSpec],
    pli_mode: PliMode,
) -> tuple[list[PliBlock], list[StageBandSpec], StageScope]:
    """Segment PLI blocks when mode is SECTION_PER_PLI; otherwise pass bands through.

    Returns (pli_blocks, sheet_stage_bands, stage_scope). For SECTION_PER_PLI the
    stage bands are consumed into block-local scope; for all other modes they remain
    at sheet level.
    """
    if pli_mode is PliMode.SECTION_PER_PLI:
        blocks = segment_blocks(rows, kv_anchors, signals.blank_run_gaps, stage_bands)
        return blocks, [], StageScope.PLI_LOCAL
    return [], stage_bands, StageScope.SHEET_LEVEL


def _compute_confidence(pli_mode: PliMode, kv_anchors: list[KVAnchor]) -> float:
    """Return a confidence score for the produced SheetPlan."""
    if pli_mode is PliMode.SHEET_IS_PLI:
        return 0.92 if len(kv_anchors) >= 3 else 0.85
    return 0.9


def _count_populated_data_cols(ws: object, plan: SheetPlan) -> int:
    """Count columns that have at least one non-empty cell across ANCHOR/CHILD data rows.

    Used to judge whether the labels collected from header_rows are sufficient.
    Returns 0 when there are no classified data rows (e.g. plan has no rows yet).
    """
    data_rows = [
        r.idx for r in plan.rows
        if r.role.value in {"anchor", "child"}
    ]
    if not data_rows:
        return 0
    cols: set[int] = set()
    for r in data_rows:
        for c_idx in range(1, (ws.max_column or 0) + 1):
            if ws.cell(row=r, column=c_idx).value is not None:
                cols.add(c_idx)
    return len(cols)


def _detect_title_row_extra_header(ws: object, plan: SheetPlan) -> int | None:
    """Return the row index of a hidden true-header row, or None.

    When the labels harvested from plan.header_rows cover fewer than half the
    populated data columns, the header_rows are pointing at a wide title row
    rather than real column labels. In that case the row immediately after
    max(header_rows) is the true header row — return its index so callers can
    both extract labels from it AND reclassify it as HEADER in plan.rows.
    """
    if not plan.header_rows or plan.pli_mode is not PliMode.ROW_PER_PLI:
        return None
    claimed: set[str] = set()
    for band in plan.stage_bands:
        for sc in band.stage_columns:
            claimed.add(sc.primary_col)
            claimed.update(sc.sub_columns.values())
    initial_labels_count = sum(
        1 for c_idx in range(1, (ws.max_column or 0) + 1)
        if get_column_letter(c_idx) not in claimed
        and any(
            isinstance(ws.cell(row=h_row, column=c_idx).value, str)
            and ws.cell(row=h_row, column=c_idx).value.strip()
            for h_row in plan.header_rows
        )
    )
    populated = _count_populated_data_cols(ws, plan)
    if populated > 0 and initial_labels_count < populated / 2:
        return max(plan.header_rows) + 1
    return None


def _collect_header_labels(ws: object, plan: SheetPlan) -> list[HeaderLabel]:
    """Lift identity-column header strings from header_rows into the artifact.

    For each column NOT claimed by a stage band's primary_col or sub_columns,
    take the first non-empty string scanning header_rows top-to-bottom. Returns
    empty for non-ROW_PER_PLI modes.

    When the title-row pattern has been detected by SheetRowPlanner.run() and the
    plan's header_rows already include the true header row, this function simply
    reads from the provided header_rows — no further detection is done here.
    """
    if plan.pli_mode is not PliMode.ROW_PER_PLI:
        return []
    claimed: set[str] = set()
    for band in plan.stage_bands:
        for sc in band.stage_columns:
            claimed.add(sc.primary_col)
            claimed.update(sc.sub_columns.values())
    labels: list[HeaderLabel] = []
    for c_idx in range(1, (ws.max_column or 0) + 1):
        col = get_column_letter(c_idx)
        if col in claimed:
            continue
        for h_row in plan.header_rows:
            v = ws.cell(row=h_row, column=c_idx).value
            if isinstance(v, str) and v.strip():
                labels.append(HeaderLabel(raw=v.strip(), col=col, row=h_row))
                break
    return labels


def _apply_extra_header_row(rows: list[RowSpec], extra_row: int) -> list[RowSpec]:
    """Reclassify a data row as HEADER when the title-row pattern is detected.

    When _detect_title_row_extra_header finds that the row immediately after the
    classified header_rows is the true header row (not a data row), reclassify
    it from ANCHOR/CHILD/BLANK to HEADER so apply_plan skips it correctly.
    Also reclassify any CHILD rows whose anchor_idx pointed to extra_row, since
    those rows are typically sub-header rows (e.g. stage sub-labels) that should
    also be HEADER rather than data.
    """
    updated = []
    for r in rows:
        if r.idx == extra_row and r.role.value in {"anchor", "child", "blank"}:
            updated.append(r.model_copy(update={
                "role": RowRole.HEADER, "anchor_idx": None, "group_id": None,
            }))
        elif r.role is RowRole.CHILD and r.anchor_idx == extra_row:
            # Sub-header row that was parented to the reclassified header row.
            updated.append(r.model_copy(update={
                "role": RowRole.HEADER, "anchor_idx": None, "group_id": None,
            }))
        else:
            updated.append(r)
    return updated


@component
class SheetRowPlanner:
    """Orchestrates deterministic detectors to produce a SheetPlan from a WorkbookCtx and a sheet name."""

    # allow-long: orchestrates detectors + title-row correction + header_labels collection
    @component.output_types(plan=SheetPlan)
    def run(self, workbook_ctx: Any, sheet: str) -> dict:
        """Run the full planner pipeline for one sheet and return a SheetPlan."""
        log.info("planner_start", sheet=sheet)

        signals = survey_sheet(workbook_ctx, sheet)
        pli_mode = _decide_pli_mode(signals)
        identity_column = (
            _pick_identity_column(signals) if pli_mode is not PliMode.SHEET_IS_PLI else None
        )
        log.info("planner_mode_decided", sheet=sheet,
                 pli_mode=pli_mode.value, identity_column=identity_column)

        stage_bands = detect_stage_bands(workbook_ctx, sheet, signals)
        kv_anchors = detect_kv_anchors(workbook_ctx, sheet, signals)

        rows = _classify_sheet_rows(workbook_ctx, sheet, signals, pli_mode, identity_column)
        header_rows = [r.idx for r in rows if r.role is RowRole.HEADER] if rows else []

        blocks, stage_bands_sheet, stage_scope = _segment_pli_blocks_if_applicable(
            rows, kv_anchors, signals, stage_bands, pli_mode
        )

        plan = SheetPlan(
            sheet=sheet,
            pli_mode=pli_mode,
            identity_column=identity_column,
            header_rows=header_rows,
            rows=rows,
            pli_blocks=blocks,
            kv_anchors=kv_anchors if pli_mode is PliMode.SHEET_IS_PLI else [],
            stage_bands=stage_bands_sheet,
            stage_scope=stage_scope,
            confidence=_compute_confidence(pli_mode, kv_anchors),
        )
        ws = workbook_ctx.wb[sheet]
        # Detect and correct the title-row pattern before collecting labels:
        # when header_rows points at a wide merged title row, the next row is
        # the true header row. Reclassify it in plan.rows and extend header_rows
        # so apply_plan skips it as data. Pass only the true header row to
        # _collect_header_labels so the title string is not included in labels.
        extra_header_row = _detect_title_row_extra_header(ws, plan)
        if extra_header_row is not None:
            corrected_rows = _apply_extra_header_row(plan.rows, extra_header_row)
            # header_rows includes the title row(s), the discovered true header row,
            # and any sub-header rows that were reclassified alongside it.
            new_header_idxs = {r.idx for r in corrected_rows if r.role is RowRole.HEADER}
            corrected_header_rows = sorted(new_header_idxs)
            plan = plan.model_copy(update={
                "rows": corrected_rows,
                "header_rows": corrected_header_rows,
            })
            # Collect labels from the true header row only (not the title row or sub-header).
            labels_plan = plan.model_copy(update={"header_rows": [extra_header_row]})
            plan = plan.model_copy(update={"header_labels": _collect_header_labels(ws, labels_plan)})
        else:
            plan = plan.model_copy(update={"header_labels": _collect_header_labels(ws, plan)})
        log.info("planner_complete", sheet=sheet, pli_mode=plan.pli_mode.value,
                 rows=len(plan.rows), blocks=len(plan.pli_blocks),
                 kv=len(plan.kv_anchors), header_labels=len(plan.header_labels),
                 bands=len(plan.stage_bands),
                 confidence=plan.confidence)
        return {"plan": plan}
