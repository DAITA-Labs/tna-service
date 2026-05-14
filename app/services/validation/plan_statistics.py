"""Tier 2 plan validators — statistical sanity of a SheetPlan against the sheet.

Warnings here cause PlanReviewer to fire; errors trigger re-plan with hints.
"""
from __future__ import annotations

from datetime import date, datetime

from openpyxl.utils import column_index_from_string

from app.core.logs import get_logger
from app.enums.row_role import RowRole
from app.enums.validation_severity import ValidationSeverity
from app.models.artifacts import SheetPlan, ValidationFinding
from app.models.workbook import WorkbookCtx

log = get_logger(__name__)


def _warn(check: str, msg: str) -> ValidationFinding:
    """Build a WARN-severity finding for the given check name and message."""
    return ValidationFinding(check=check, severity=ValidationSeverity.WARN, message=msg)


def _check_sequence_numbers(ctx: WorkbookCtx, plan: SheetPlan) -> list[ValidationFinding]:
    """Verify S.NO column max matches anchor count; emit a finding when they diverge."""
    ws = ctx.wb[plan.sheet]
    anchors = [r for r in plan.rows if r.role is RowRole.ANCHOR]
    sno_col = None
    for col in range(1, (ws.max_column or 0) + 1):
        for h_row in plan.header_rows:
            cell_val = ws.cell(row=h_row, column=col).value
            if isinstance(cell_val, str) and cell_val.strip().lower() in (
                "s no", "s.no", "sno", "s. no", "sl no"
            ):
                sno_col = col
                break
        if sno_col:
            break
    if sno_col is None or not anchors:
        return []
    sno_vals = [ws.cell(row=a.idx, column=sno_col).value for a in anchors]
    sno_ints = [int(v) for v in sno_vals if isinstance(v, (int, float))]
    if sno_ints and max(sno_ints) != len(anchors):
        return [_warn("sequence_match",
                      f"S.NO max={max(sno_ints)} vs ANCHOR count={len(anchors)}")]
    return []


def _check_pli_count_sanity(ctx: WorkbookCtx, plan: SheetPlan) -> list[ValidationFinding]:
    """Emit a finding when a non-trivial sheet produces no PLIs, blocks, or KV anchors."""
    ws = ctx.wb[plan.sheet]
    max_row = ws.max_row or 0
    anchors = [r for r in plan.rows if r.role is RowRole.ANCHOR]
    if (max_row >= 10
            and len(anchors) == 0
            and len(plan.pli_blocks) == 0
            and len(plan.kv_anchors) == 0):
        return [_warn("pli_count_sanity",
                      f"sheet has {max_row} rows but plan emits 0 PLIs")]
    return []


def _check_identity_column_coverage(ctx: WorkbookCtx, plan: SheetPlan) -> list[ValidationFinding]:
    """Emit a finding when the identity column is sparsely populated across data rows."""
    if not plan.identity_column or not plan.rows:
        return []
    ws = ctx.wb[plan.sheet]
    id_col = column_index_from_string(plan.identity_column)
    data_rows = [r for r in plan.rows if r.role in (RowRole.ANCHOR, RowRole.CHILD)]
    if not data_rows:
        return []
    populated = sum(
        1 for r in data_rows
        if ws.cell(row=r.idx, column=id_col).value is not None
        or r.anchor_idx is not None
    )
    ratio = populated / len(data_rows)
    if ratio < 0.8:
        return [_warn("identity_column_coverage",
                      f"identity coverage {ratio:.2f} < 0.80")]
    return []


def _check_date_band_density(ctx: WorkbookCtx, plan: SheetPlan) -> list[ValidationFinding]:
    """Emit a finding for any stage band whose stage-column cells are mostly non-date values."""
    ws = ctx.wb[plan.sheet]
    findings: list[ValidationFinding] = []
    for band in plan.stage_bands:
        sub_rows = list(band.sub_rows.values())
        if not sub_rows:
            continue
        total = 0
        dates = 0
        for col_letter in band.stage_cols.values():
            col_idx = column_index_from_string(col_letter)
            for row in sub_rows:
                val = ws.cell(row=row, column=col_idx).value
                if val is None:
                    continue
                total += 1
                if isinstance(val, (date, datetime)):
                    dates += 1
        if total >= 2 and dates / total < 0.5:
            findings.append(_warn("date_band_density",
                                  f"stage band '{band.name}' is only {dates}/{total} date-typed"))
    return findings


def validate_statistics(ctx: WorkbookCtx, plan: SheetPlan) -> list[ValidationFinding]:
    """Run all Tier 2 statistical sanity checks against a SheetPlan.

    Returns the concatenated findings. WARN severity causes PlanReviewer to fire.
    """
    findings = [
        *_check_sequence_numbers(ctx, plan),
        *_check_pli_count_sanity(ctx, plan),
        *_check_identity_column_coverage(ctx, plan),
        *_check_date_band_density(ctx, plan),
    ]
    if findings:
        log.info("plan_statistics_warns", sheet=plan.sheet, count=len(findings),
                 checks=[f.check for f in findings])
    return findings
