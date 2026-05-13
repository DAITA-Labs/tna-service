"""Deterministic stage applier.

Reads stage planned_date + sub_column values per PLI row. Supports both
wide_sub_columns (Plan/Actual side-by-side) and tall_sub_rows
(Plan/Action/Deviation stacked). Merge propagation kicks in for vertical_merge
layouts so multi-color sub-rows inherit stage dates from the merge anchor.

`strip_stage_columns_from_metadata` deduplicates: any FieldMap.metadata_locations
column that's already claimed by a StageColumn.primary_col or sub_columns gets
dropped — the orchestrator calls this before invoking the field applier.
"""
from __future__ import annotations
import logging
from typing import Any
from openpyxl.utils import column_index_from_string, get_column_letter
from pydantic import ValidationError
from app.models.extraction import Stage
from app.models.workbook import WorkbookCtx
from app.models.artifacts import (
    StageBandSet, StageBand, StageColumn, PLIBoundaries, FieldMap,
)
from app.enums.boundary_pattern import BoundaryPattern
from app.enums.stage_layout_mode import StageLayoutMode
from app.services.applier._registry import get_pattern_handler

log = logging.getLogger(__name__)


def strip_stage_columns_from_metadata(fm: FieldMap, sset: StageBandSet) -> FieldMap:
    """Drop metadata_locations whose column is claimed by any stage's
    primary_col or sub_columns."""
    claimed: set[str] = set()
    for band in sset.bands:
        for sc in band.stage_columns:
            if sc.primary_col:
                claimed.add(sc.primary_col)
            for col in sc.sub_columns.values():
                if col:
                    claimed.add(col)
    if not claimed:
        return fm
    kept = [m for m in fm.metadata_locations if m.column not in claimed]
    if len(kept) == len(fm.metadata_locations):
        return fm
    return fm.model_copy(update={"metadata_locations": kept})


def _build_stage_tolerantly(values: dict) -> Stage:
    """Drop fields that fail validation; preserve name."""
    attempt = dict(values)
    for _ in range(len(attempt) + 1):
        try:
            return Stage(**attempt)
        except ValidationError as e:
            dropped = [
                err["loc"][0] for err in e.errors()
                if err["loc"] and isinstance(err["loc"][0], str)
                and err["loc"][0] in attempt and err["loc"][0] != "name"
            ]
            if not dropped:
                raise
            for k in dropped:
                log.warning("Stage %r dropped %r: %s",
                            attempt.get("name"), k, e.errors()[0].get("msg", ""))
                attempt.pop(k, None)
    return Stage(name=values.get("name", "<unnamed>"))


def _read_cell(ws, row: int, col_idx: int, propagate: bool) -> tuple[Any, str]:
    """Return (value, source A1 address) — with merge propagation when asked."""
    if propagate:
        for mr in ws.merged_cells.ranges:
            if (mr.min_row <= row <= mr.max_row
                    and mr.min_col <= col_idx <= mr.max_col):
                addr = f"{get_column_letter(mr.min_col)}{mr.min_row}"
                return ws.cell(row=mr.min_row, column=mr.min_col).value, addr
    return (ws.cell(row=row, column=col_idx).value,
            f"{get_column_letter(col_idx)}{row}")


def _read_stage_wide(ws, band: StageBand, sc: StageColumn,
                     pli_row: int, propagate: bool) -> Stage:
    primary_idx = column_index_from_string(sc.primary_col)
    planned, planned_addr = _read_cell(ws, pli_row, primary_idx, propagate)
    metadata: dict[str, Any] = {}
    source_cells: dict[str, str] = {"planned_date": planned_addr}
    for key, col_letter in sc.sub_columns.items():
        v, addr = _read_cell(ws, pli_row, column_index_from_string(col_letter), propagate)
        if v is not None:
            metadata[key] = v
            source_cells[key] = addr
    return _build_stage_tolerantly(dict(
        name=sc.name, planned_date=planned, section=band.section_name,
        metadata=metadata, source_cells=source_cells, confidence=band.confidence,
    ))


def _read_stage_tall(ws, band: StageBand, sc: StageColumn) -> Stage:
    primary_idx = column_index_from_string(sc.primary_col)
    primary_letter = get_column_letter(primary_idx)
    plan_row = (band.sub_rows.get("Plan") or band.sub_rows.get("plan")
                or band.name_row + 1)
    planned = ws.cell(row=plan_row, column=primary_idx).value
    metadata: dict[str, Any] = {}
    source_cells: dict[str, str] = {"planned_date": f"{primary_letter}{plan_row}"}
    for key, row_idx in band.sub_rows.items():
        if key.lower() == "plan":
            continue
        v = ws.cell(row=row_idx, column=primary_idx).value
        if v is not None:
            clean = key.lower().replace(" ", "_").replace("if_any", "").rstrip("_")
            metadata[clean] = v
            source_cells[clean] = f"{primary_letter}{row_idx}"
    return _build_stage_tolerantly(dict(
        name=sc.name, planned_date=planned, section=band.section_name,
        metadata=metadata, source_cells=source_cells, confidence=band.confidence,
    ))


def apply_stage_band_set(
    ctx: WorkbookCtx, sheet: str,
    sset: StageBandSet, boundaries: PLIBoundaries,
) -> list[list[Stage]]:
    """Return list[list[Stage]] — one inner list per PLI row (parallel to apply_field_map)."""
    if boundaries.pattern == BoundaryPattern.ONE_SHEET_PER_PLI:
        out: list[list[Stage]] = []
        for s in boundaries.sheet_iter:
            ws = ctx.wb[s]
            stages: list[Stage] = []
            for band in sset.bands:
                for sc in band.stage_columns:
                    if band.layout_mode == StageLayoutMode.TALL_SUB_ROWS:
                        stages.append(_read_stage_tall(ws, band, sc))
                    else:
                        row = band.data_start_row or band.name_row + 1
                        stages.append(_read_stage_wide(ws, band, sc, row, False))
            out.append(stages)
        return out

    pattern_value = (boundaries.pattern.value
                     if hasattr(boundaries.pattern, "value") else boundaries.pattern)
    handler = get_pattern_handler(pattern_value)
    pli_rows = handler(ctx, boundaries)
    propagate = boundaries.pattern == BoundaryPattern.VERTICAL_MERGE
    ws = ctx.wb[sheet]
    out = []
    for r in pli_rows:
        stages: list[Stage] = []
        for band in sset.bands:
            for sc in band.stage_columns:
                if band.layout_mode == StageLayoutMode.WIDE_SUB_COLUMNS:
                    stages.append(_read_stage_wide(ws, band, sc, r, propagate))
                else:
                    stages.append(_read_stage_tall(ws, band, sc))
        out.append(stages)
    return out
