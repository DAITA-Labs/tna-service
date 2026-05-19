"""Deterministic plan executor: (SheetPlan, CanonicalNameMap, WorkbookCtx) -> list[PLI].

Receives a `SheetPlan` produced by the SheetRowPlanner pipeline and a
`CanonicalNameMap` built from supplier-specific label normalisations, then
iterates the workbook cells to produce fully-populated `PLI` objects.
Dispatches on `pli_mode` (ROW_PER_PLI / SECTION_PER_PLI / SHEET_IS_PLI) —
each mode has its own private apply function; `apply_plan` is the single public
entry point.  This module is statically guaranteed to import no LLM modules; the
test `tests/unit/applier/test_apply_plan_no_llm_imports.py` enforces this via
an AST import scan.
"""
from __future__ import annotations

from datetime import date, datetime

from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string

from app.core.logs import get_logger
from app.enums.pli_mode import PliMode
from app.enums.row_role import RowRole
from app.models.artifacts import (
    CanonicalNameMap,
    KVAnchor,
    PliBlock,
    RowSpec,
    SheetPlan,
    StageBandSpec,
)
from app.models.extraction import PLI, Stage
from app.models.workbook import WorkbookCtx

log = get_logger(__name__)


_DATA_ROLES = {RowRole.ANCHOR, RowRole.CHILD}
_STRING_FIELDS = {"io_number", "style_code", "style_name",
                  "color_code", "color_name", "fabric_code"}

_CONFIDENCE_DEFAULTS = {
    "kv_anchor": 0.95,
    "header_label": 0.85,
    "stage_column": 0.85,
    "stage_subfield": 0.80,
    "metadata_fallback": 0.40,
}


def _resolve_confidence(*, source: str, name_map: CanonicalNameMap,
                        raw: str, canonical: str) -> float:
    """Pick a per-field confidence value.

    LLM-supplied confidence wins when the canonical name appears in
    `name_map.field_confidence`. Otherwise returns a calibrated default per
    source type. Returns 0.5 for unknown sources.
    """
    if canonical in name_map.field_confidence:
        return name_map.field_confidence[canonical]
    return _CONFIDENCE_DEFAULTS.get(source, 0.5)


def _coerce(field: str, val: object) -> object:
    """Coerce a raw cell value to the type expected by `field`.

    String-typed PLI fields (style_code, color_code, etc.) are stringified even
    when the cell holds an int — Excel sometimes stores codes as numbers.
    """
    if field in _STRING_FIELDS and val is not None:
        return str(val)
    return val


def _read_with_merge(ws, row: int, col_idx: int) -> tuple[object, str]:
    """Return (value, cell-address) for a cell, following merged-cell anchors.

    If the cell is empty but falls inside a merged region, returns the value and
    address of the top-left anchor cell of that region.  Returns (None, addr) when
    the cell and any enclosing merge are both empty.
    """
    direct = ws.cell(row=row, column=col_idx).value
    if direct is not None:
        return direct, f"{get_column_letter(col_idx)}{row}"
    for mr in ws.merged_cells.ranges:
        if mr.min_row <= row <= mr.max_row and mr.min_col <= col_idx <= mr.max_col:
            anc_val = ws.cell(row=mr.min_row, column=mr.min_col).value
            return anc_val, f"{get_column_letter(mr.min_col)}{mr.min_row}"
    return None, f"{get_column_letter(col_idx)}{row}"


def _read_kv_into(values: dict, source_cells: dict, ws, kv: KVAnchor,
                  name_map: CanonicalNameMap) -> None:
    """Read one KV anchor cell and write the result into `values` and `source_cells`.

    Mutates both dicts in place.  Skips the field when the canonical name is
    "ignore" or when the cell is empty.
    """
    canonical = name_map.field_labels.get(kv.field, kv.field)
    if canonical == "ignore":
        return
    col_letter, row = coordinate_from_string(kv.value_cell)
    col = column_index_from_string(col_letter)
    val = ws.cell(row=row, column=col).value
    if val is None:
        return
    values[canonical] = _coerce(canonical, val)
    source_cells[canonical] = kv.value_cell


def _read_wide_stage_column(ws, band: StageBandSpec, stage_name: str,
                            col_letter: str, pli_row: int,
                            name_map: CanonicalNameMap) -> Stage | None:
    """Read one stage column in wide_sub_columns layout for `pli_row`.

    Returns a `Stage` when the cell is non-empty, or None to signal skip.
    """
    canonical_stage = name_map.stage_names.get(stage_name, stage_name)
    if canonical_stage == "ignore":
        return None
    c_idx = column_index_from_string(col_letter)
    val, addr = _read_with_merge(ws, pli_row, c_idx)
    if val is None:
        return None
    return Stage(
        name=canonical_stage,
        planned_date=val if isinstance(val, (date, datetime)) else None,
        section=band.name,
        source={"sheet": ws.title, "rows": [pli_row],
                "cells": {"planned_date": addr}},
    )


def _read_tall_stage_column(ws, band: StageBandSpec, stage_name: str,
                            col_letter: str,
                            name_map: CanonicalNameMap) -> Stage | None:
    """Read one stage column in tall_sub_rows layout using band.sub_rows.

    Reads the "plan" sub-row for the planned date and every other sub-row into
    `metadata`.  Returns None when the plan row is absent or its cell is empty.
    """
    canonical_stage = name_map.stage_names.get(stage_name, stage_name)
    if canonical_stage == "ignore":
        return None
    plan_row = band.sub_rows.get("plan")
    if plan_row is None:
        return None
    c_idx = column_index_from_string(col_letter)
    pv, pa = _read_with_merge(ws, plan_row, c_idx)
    if pv is None:
        return None
    metadata: dict = {}
    source_cells = {"planned_date": pa}
    for role, row_idx in band.sub_rows.items():
        if role == "plan":
            continue
        av = ws.cell(row=row_idx, column=c_idx).value
        if av is not None:
            metadata[role] = av
            source_cells[role] = f"{col_letter}{row_idx}"
    return Stage(
        name=canonical_stage,
        planned_date=pv if isinstance(pv, (date, datetime)) else None,
        section=band.name, metadata=metadata,
        source={"sheet": ws.title, "rows": sorted(band.sub_rows.values()),
                "cells": source_cells},
    )


def _read_stages(ws, bands: list[StageBandSpec], pli_row: int,
                 name_map: CanonicalNameMap, parent_source: dict) -> list[Stage]:
    """Collect all Stage objects for a PLI row across every stage band.

    Dispatches each stage column to `_read_wide_stage_column` or
    `_read_tall_stage_column` depending on `band.layout_mode`.
    `parent_source` is accepted for interface symmetry but is not read here.
    """
    stages: list[Stage] = []
    for band in bands:
        for stage_name, col_letter in band.stage_cols.items():
            if band.layout_mode == "wide_sub_columns":
                stage = _read_wide_stage_column(
                    ws, band, stage_name, col_letter, pli_row, name_map)
            else:  # tall_sub_rows
                stage = _read_tall_stage_column(
                    ws, band, stage_name, col_letter, name_map)
            if stage is not None:
                stages.append(stage)
    return stages


def _emit_single_row_pli(ws, plan: SheetPlan, row: RowSpec,
                         header_label_by_col: dict[int, str],
                         name_map: CanonicalNameMap) -> PLI:
    """Build one PLI from a single data row plus plan-level KV anchors and stages."""
    values: dict = {"metadata": {}}
    source_cells: dict[str, str] = {}

    for col_idx, label in header_label_by_col.items():
        canonical = name_map.field_labels.get(label, label)
        if canonical == "ignore":
            continue
        val, addr = _read_with_merge(ws, row.idx, col_idx)
        if val is None:
            continue
        if canonical in PLI.model_fields:
            values[canonical] = _coerce(canonical, val)
            source_cells[canonical] = addr
        else:
            values["metadata"][canonical] = val
            source_cells[canonical] = addr

    for kv in plan.kv_anchors:
        _read_kv_into(values, source_cells, ws, kv, name_map)

    values["source"] = {"sheet": plan.sheet, "rows": [row.idx], "cells": source_cells}
    values["stages"] = _read_stages(ws, plan.stage_bands, row.idx, name_map, source_cells)
    return PLI(**values)


def _emit_multi_row_pli(ws, plan: SheetPlan, group_rows: list[RowSpec],
                        header_label_by_col: dict[int, str],
                        name_map: CanonicalNameMap) -> PLI:
    """Build one PLI from a multi-row group, using the ANCHOR row as the primary source.

    Sub-row indices and roles are recorded in `pli.metadata["multi_row_sub_rows"]`.
    """
    anchor_row = next(r for r in group_rows if r.role is RowRole.ANCHOR)
    pli = _emit_single_row_pli(ws, plan, anchor_row, header_label_by_col, name_map)
    for r in group_rows:
        if r is anchor_row:
            continue
        pli.metadata.setdefault("multi_row_sub_rows", []).append({
            "row": r.idx, "sub_row_role": r.sub_row_role.value if r.sub_row_role else None,
        })
    return pli


def _apply_row_per_pli(ctx: WorkbookCtx, plan: SheetPlan,
                       name_map: CanonicalNameMap) -> list[PLI]:
    """Produce one PLI per data row (or per grouped multi-row block) in the sheet."""
    ws = ctx.wb[plan.sheet]
    plis: list[PLI] = []
    rows_by_group: dict[int, list[RowSpec]] = {}
    for r in plan.rows:
        if r.role not in _DATA_ROLES:
            continue
        if r.group_id is not None:
            rows_by_group.setdefault(r.group_id, []).append(r)
        else:
            rows_by_group.setdefault(r.idx, []).append(r)

    header_label_by_col: dict[int, str] = {}
    for h_row in plan.header_rows:
        for c in range(1, (ws.max_column or 0) + 1):
            v = ws.cell(row=h_row, column=c).value
            if isinstance(v, str) and v.strip():
                header_label_by_col.setdefault(c, v.strip())

    for gid, group_rows in rows_by_group.items():
        is_multi_row_pli = all(r.sub_row_role is not None for r in group_rows)
        if is_multi_row_pli:
            plis.append(_emit_multi_row_pli(ws, plan, group_rows,
                                            header_label_by_col, name_map))
            continue
        for r in group_rows:
            plis.append(_emit_single_row_pli(ws, plan, r,
                                             header_label_by_col, name_map))
    return plis


def _apply_section_per_pli(ctx: WorkbookCtx, plan: SheetPlan,
                           name_map: CanonicalNameMap) -> list[PLI]:
    """Produce one PLI per `PliBlock` — each block is a vertical KV section."""
    ws = ctx.wb[plan.sheet]
    plis: list[PLI] = []
    for blk in plan.pli_blocks:
        values: dict = {"metadata": {}}
        source_cells: dict[str, str] = {}
        for kv in blk.identity:
            _read_kv_into(values, source_cells, ws, kv, name_map)
        values["source"] = {"sheet": plan.sheet,
                            "rows": list(range(blk.bbox[0], blk.bbox[1] + 1)),
                            "cells": source_cells}
        values["stages"] = []
        for band in blk.stage_bands:
            plan_row = band.sub_rows.get("plan")
            if plan_row is not None:
                values["stages"].extend(
                    _read_stages(ws, [band], plan_row, name_map, source_cells)
                )
        plis.append(PLI(**values))
    return plis


def _apply_sheet_is_pli(ctx: WorkbookCtx, plan: SheetPlan,
                        name_map: CanonicalNameMap) -> list[PLI]:
    """Produce a single PLI for the entire sheet — the sheet itself is one product line."""
    ws = ctx.wb[plan.sheet]
    values: dict = {"metadata": {}}
    source_cells: dict[str, str] = {}
    for kv in plan.kv_anchors:
        _read_kv_into(values, source_cells, ws, kv, name_map)
    values["source"] = {"sheet": plan.sheet, "rows": [], "cells": source_cells}
    values["stages"] = []
    for band in plan.stage_bands:
        plan_row = band.sub_rows.get("plan")
        if plan_row is not None:
            values["stages"].extend(
                _read_stages(ws, [band], plan_row, name_map, source_cells)
            )
    return [PLI(**values)]


def apply_plan(ctx: WorkbookCtx, plan: SheetPlan,
               name_map: CanonicalNameMap) -> list[PLI]:
    """Dispatch on pli_mode and return all PLIs extracted from `plan`.

    Raises `ValueError` for unknown `pli_mode` values.  Pure function — no LLM
    calls, no I/O beyond reading the workbook already loaded in `ctx`.
    """
    log.info("apply_plan_start", sheet=plan.sheet, pli_mode=plan.pli_mode.value)
    if plan.pli_mode is PliMode.ROW_PER_PLI:
        result = _apply_row_per_pli(ctx, plan, name_map)
    elif plan.pli_mode is PliMode.SECTION_PER_PLI:
        result = _apply_section_per_pli(ctx, plan, name_map)
    elif plan.pli_mode is PliMode.SHEET_IS_PLI:
        result = _apply_sheet_is_pli(ctx, plan, name_map)
    else:
        raise ValueError(f"unknown pli_mode: {plan.pli_mode}")
    log.info("apply_plan_complete", sheet=plan.sheet, pli_count=len(result))
    return result
