"""Search tools — locate a value within a sheet."""
from __future__ import annotations
from openpyxl.utils import get_column_letter
from app.models.workbook import WorkbookCtx
from app.tools._decorator import tool


@tool("find_value")
def find_value(
    ctx: WorkbookCtx, sheet: str, needle: str,
    max_hits: int = 10, case_insensitive: bool = True,
) -> list[str]:
    """Return A1 addresses of cells containing `needle` (substring match)."""
    ws = ctx.wb[sheet]
    target = needle.lower() if case_insensitive else needle
    hits: list[str] = []
    for r in range(1, (ws.max_row or 0) + 1):
        for c in range(1, (ws.max_column or 0) + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            s = str(v).lower() if case_insensitive else str(v)
            if target in s:
                hits.append(f"{get_column_letter(c)}{r}")
                if len(hits) >= max_hits:
                    return hits
    return hits
