"""KvBlock detector — bold/filled label cells paired with an adjacent value cell.

A KvBlock is the structural signature of SHEET_IS_PLI / scattered-metadata
layouts: a label cell that's visually distinct (bold or filled) sitting
next to a value cell that isn't itself a label.

Examples (63261-TNA):
    A4 = "Job No"    B4 = 63261      → KvBlock(A4 → B4)
    A5 = "Quantity"  B5 = 16200      → KvBlock(A5 → B5)
    D4 = "Ex-Fty"    E4 = 2026-05-07 → KvBlock(D4 → E4)

KvBlockDetector does NOT use spec aliases — it identifies the visual
pattern only. Downstream components (IoNumberComponent, QuantityComponent,
date trio) consume KvBlock candidates and match the label text against
their own spec aliases.

Search order: right → below → above → left. This matches the typical
k:v layout where the value sits to the right of the label, with column
or above/left fallbacks for stacked metadata blocks.
"""
from __future__ import annotations

from openpyxl.utils import get_column_letter

from app.artifacts.canvas import GridCanvas
from app.artifacts.structure import KvBlock
from app.tools._decorator import tool


_DTYPE_BLANK = 0


@tool("find_kv_blocks")
def find_kv_blocks(canvas: GridCanvas) -> list[KvBlock]:
    """Detect every KvBlock on the canvas.

    A cell qualifies as a label if it's bold OR has a non-zero fill colour
    AND its string content is non-empty. For each label, search adjacent
    cells (right → below → above → left) for a non-blank value cell whose
    own visual styling is NOT label-like (not bold, not filled). The first
    match wins.
    """
    bold = canvas.channels.get("bold")
    fill = canvas.channels.get("fill_color")
    dtype = canvas.channels.get("dtype")
    if bold is None or fill is None or dtype is None:
        return []

    label_positions = _label_candidate_positions(canvas, bold, fill)
    if not label_positions:
        return []

    blocks: list[KvBlock] = []
    for r, c in label_positions:
        value_pos = _find_adjacent_value(canvas, r, c, label_positions, bold, fill, dtype)
        if value_pos is None:
            continue
        vr, vc = value_pos
        blocks.append(KvBlock(
            label_coord=(get_column_letter(c + 1), r + 1),
            value_coord=(get_column_letter(vc + 1), vr + 1),
            label_text=str(canvas.cell_values[r][c]).strip(),
            value_dtype=int(dtype[vr][vc]),
        ))
    return blocks


# ─── Helpers ────────────────────────────────────────────────────────────────


def _label_candidate_positions(canvas: GridCanvas,
                                bold: list[list[int]],
                                fill: list[list[int]]) -> set[tuple[int, int]]:
    """Return 0-indexed (row, col) positions of bold OR filled cells with string content."""
    positions: set[tuple[int, int]] = set()
    for r in range(canvas.n_rows):
        for c in range(canvas.n_cols):
            is_label_styled = (bold[r][c] == 1) or (fill[r][c] > 0)
            if not is_label_styled:
                continue
            value = canvas.cell_values[r][c]
            if not isinstance(value, str) or not value.strip():
                continue
            positions.add((r, c))
    return positions


def _find_adjacent_value(canvas: GridCanvas,
                          r: int, c: int,
                          label_positions: set[tuple[int, int]],
                          bold: list[list[int]],
                          fill: list[list[int]],
                          dtype: list[list[int]]) -> tuple[int, int] | None:
    """Search right → below → above → left for a non-label, non-blank neighbour."""
    for dr, dc in ((0, 1), (1, 0), (-1, 0), (0, -1)):
        nr, nc = r + dr, c + dc
        if not (0 <= nr < canvas.n_rows and 0 <= nc < canvas.n_cols):
            continue
        if (nr, nc) in label_positions:
            continue
        if dtype[nr][nc] == _DTYPE_BLANK:
            continue
        # The neighbour must NOT itself look like a label
        if bold[nr][nc] == 1 or fill[nr][nc] > 0:
            continue
        return (nr, nc)
    return None
