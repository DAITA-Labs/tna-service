"""DKN-shape minimal fixture: two-row stage header (name row + sub-label row)."""
from datetime import date

from openpyxl import Workbook


def build(wb: Workbook) -> None:
    """Populate a wb with a DKN-style two-row stage header layout."""
    ws = wb.active
    ws.title = "DKN"
    # Row 1 identity headers
    ws["B1"] = "Buyer Po No"
    ws["C1"] = "Color"
    ws["D1"] = "Order Qty"
    ws["E1"] = "Etd Ex factory as per P.O"
    # Row 1 stage names merged across each stage's column block
    # CUTTING spans cols H..I (Plan, Actual); SEWING spans J..K
    ws["H1"] = "CUTTING"
    ws.merge_cells("H1:I1")
    ws["J1"] = "SEWING"
    ws.merge_cells("J1:K1")
    # Row 2 sub-labels under each stage
    ws["H2"] = "Plan"
    ws["I2"] = "Actual"
    ws["J2"] = "Plan"
    ws["K2"] = "Actual"
    # Data row 3
    ws["B3"] = "PO-00045"
    ws["C3"] = "Navy"
    ws["D3"] = 500
    ws["E3"] = date(2026, 5, 5)
    ws["H3"] = date(2026, 3, 12)
    ws["I3"] = date(2026, 3, 16)
    ws["J3"] = date(2026, 4, 3)
    ws["K3"] = date(2026, 4, 5)
