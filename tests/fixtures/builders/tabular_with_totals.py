"""Tabular layout with merged identity per IO and an interleaved total row.

Family 2 anchor — CHRISTIAN BERG-style at a minimal scale.
2 IO groups: IO 1063 has 2 colors, IO 1064 has 1 color, plus 2 total rows.
"""
from openpyxl import Workbook


def build(wb: Workbook) -> None:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    ws = wb.create_sheet("S")
    ws["A1"] = "S NO"; ws["B1"] = "IO NO"; ws["K1"] = "COLOR"; ws["L1"] = "ORDER QTY"
    ws["A2"] = 1; ws["B2"] = "1063"; ws["K2"] = "MAGENTA"; ws["L2"] = 1000
    ws["K3"] = "NAVY"; ws["L3"] = 1000
    ws["L4"] = 2000              # total for group 0
    ws["A5"] = 2; ws["B5"] = "1064"; ws["K5"] = "PINE"; ws["L5"] = 1500
    ws["L6"] = 1500              # total for group 1
    ws.merge_cells("A2:A3"); ws.merge_cells("B2:B3")
