"""Flat tabular layout: header row + 3 data rows, no merges, no totals.

Family 1 anchor — GUESS ATHLEISURE-style.
"""
from openpyxl import Workbook


def build(wb: Workbook) -> None:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    ws = wb.create_sheet("S")
    ws["A1"] = "IO NO"
    ws["B1"] = "STYLE"
    ws["C1"] = "ORDER QTY"
    ws["A2"] = "1063"; ws["B2"] = "DWJE"; ws["C2"] = 2000
    ws["A3"] = "1064"; ws["B3"] = "DWJF"; ws["C3"] = 1500
    ws["A4"] = "1065"; ws["B4"] = "DWJG"; ws["C4"] = 3000
