"""Vertical-merge layout: 2 IOs, each with 2-3 color child rows.

Family 2 anchor — Compass Pro / MAIN FALL KIDS style.
"""
from openpyxl import Workbook


def build(wb: Workbook) -> None:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    ws = wb.create_sheet("S")
    ws["A1"] = "S NO"; ws["B1"] = "IO NO"; ws["K1"] = "COLOR"; ws["L1"] = "ORDER QTY"
    ws["A2"] = 1; ws["B2"] = "1063"; ws["K2"] = "MAGENTA"; ws["L2"] = 1000
    ws["K3"] = "NAVY"; ws["L3"] = 1000
    ws["K4"] = "WHITE"; ws["L4"] = 1000
    ws["A5"] = 2; ws["B5"] = "1064"; ws["K5"] = "PINE"; ws["L5"] = 1500
    ws["K6"] = "GRAY"; ws["L6"] = 1500
    ws.merge_cells("A2:A4"); ws.merge_cells("B2:B4")
    ws.merge_cells("A5:A6"); ws.merge_cells("B5:B6")
