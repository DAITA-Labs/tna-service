"""Tabular sheet whose header row has no IO/JOB/PO column.

Triggers planner ambiguity: no identity_col_candidates.
"""
from openpyxl import Workbook


def build(wb: Workbook) -> None:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    ws = wb.create_sheet("S")
    ws["A1"] = "NAME"; ws["B1"] = "DESCRIPTION"; ws["C1"] = "NOTES"
    ws["A2"] = "thing1"; ws["B2"] = "first thing"; ws["C2"] = "no notes"
    ws["A3"] = "thing2"; ws["B3"] = "second thing"; ws["C3"] = "more notes"
