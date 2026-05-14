"""Sheet with a stage band whose columns hold mostly strings, not dates.

Triggers Tier 2 date_band_density warning.
"""
from datetime import datetime
from openpyxl import Workbook


def build(wb: Workbook) -> None:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    ws = wb.create_sheet("S")
    ws["A1"] = "IO NO"; ws["B1"] = "STYLE"
    ws["A2"] = "1063"; ws["B2"] = "DWJE"
    ws["A3"] = "1064"; ws["B3"] = "DWJF"
    # Stage band with mostly non-date values
    ws["C5"] = "Pre-Prod"
    ws["D5"] = "L/D send"; ws["E5"] = "Fit send"
    # Row 6 has two date cells to trigger detection (acts as sub_header_row)
    ws["D6"] = datetime(2026, 3, 5); ws["E6"] = datetime(2026, 3, 5)
    # Add sub-row labels to trigger tall_sub_rows detection
    ws["A7"] = "Plan"; ws["D7"] = "pending"; ws["E7"] = "tbd"
    ws["A8"] = "Action"; ws["D8"] = "pending"; ws["E8"] = "pending"
    ws["A9"] = "Actual"; ws["D9"] = "approved"; ws["E9"] = "pending"
    # With sub_rows being [7, 8, 9], total cells = 3 rows * 2 cols = 6, dates = 0
    # Density = 0/6 = 0.0 < 0.5 triggers warning
