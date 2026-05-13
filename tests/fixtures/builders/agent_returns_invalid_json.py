"""A normal tabular sheet. The agent test will pass a FakeLLM that raises
when asked for CanonicalNameMap, simulating an agent returning unparseable output.
"""
from openpyxl import Workbook


def build(wb: Workbook) -> None:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    ws = wb.create_sheet("S")
    ws["A1"] = "IO NO"
    ws["A2"] = "1063"
