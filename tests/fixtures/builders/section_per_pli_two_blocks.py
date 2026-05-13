"""Two PLI blocks separated by a blank-row gap. Each block has its own
KV identity + a local stage band."""
from datetime import datetime
from openpyxl import Workbook


def build(wb: Workbook) -> None:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    ws = wb.create_sheet("S")
    ws["A4"] = "IO"; ws["B4"] = "1063"
    ws["C8"] = "Cut"; ws["D8"] = "Sew"
    ws["C9"] = datetime(2026, 3, 12); ws["D9"] = datetime(2026, 4, 3)
    ws["A14"] = "IO"; ws["B14"] = "1064"
    ws["C18"] = "Cut"; ws["D18"] = "Sew"
    ws["C19"] = datetime(2026, 4, 16); ws["D19"] = datetime(2026, 4, 25)
