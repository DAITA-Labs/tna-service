"""Tabular layout with a repeated header row mid-data.

Family 4 anchor — NORTHERN REFLECTIONS-style.
"""
from openpyxl import Workbook


def build(wb: Workbook) -> None:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    ws = wb.create_sheet("S")
    ws["A1"] = "IO NO"; ws["B1"] = "COLOR"
    ws["A2"] = "1063"; ws["B2"] = "MAGENTA"
    ws["A3"] = "1064"; ws["B3"] = "NAVY"
    ws["A4"] = "IO NO"; ws["B4"] = "COLOR"     # repeat header
    ws["A5"] = "1065"; ws["B5"] = "PINE"
