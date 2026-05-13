"""3-sheet workbook where each sheet is a single PLI with KV identity.

Family 5 anchor — new job-TNA.xlsx style.
"""
from datetime import datetime
from openpyxl import Workbook


def build(wb: Workbook) -> None:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    for job in ["63315", "63306", "63305"]:
        ws = wb.create_sheet(job)
        ws["A4"] = "Job No"; ws["B4"] = int(job)
        ws["A5"] = "Quantity"; ws["B5"] = 100000
        ws["A8"] = "Pre-Prod TNA"
        ws["C8"] = "L/D send"; ws["D8"] = "Fit send"
        ws["C9"] = datetime(2026, 3, 1)
        ws["D9"] = datetime(2026, 3, 5)
