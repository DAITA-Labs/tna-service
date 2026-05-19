"""Christian Berg-shape minimal fixture: single-row stage strip, ROW_PER_PLI."""
from datetime import date

from openpyxl import Workbook


def build(wb: Workbook) -> None:
    """Populate a wb with a CB-style single-row stage strip layout."""
    ws = wb.active
    ws.title = "CB"
    # Header row 3
    ws["A3"] = "S NO"
    ws["B3"] = "IO NO"
    ws["F3"] = "STYLE"
    ws["K3"] = "COLOR"
    ws["L3"] = "ORDER QTY"
    ws["D3"] = "EX FAC DATE"
    ws["N3"] = "FABRIC"
    ws["U3"] = "CUTTING"
    ws["V3"] = "Actual"
    ws["AE3"] = "SEWING"
    # Row 4
    ws["A4"] = 1
    ws["B4"] = 1063
    ws["F4"] = "T-STYLE"
    ws["K4"] = "422-MAG"
    ws["L4"] = 2356
    ws["D4"] = date(2026, 5, 5)
    ws["N4"] = date(2026, 3, 18)
    ws["U4"] = date(2026, 3, 12)
    ws["V4"] = date(2026, 3, 16)
    ws["AE4"] = date(2026, 4, 3)
    # Row 5
    ws["A5"] = 2
    ws["B5"] = 1064
    ws["F5"] = "T-STYLE2"
    ws["K5"] = "630-NAV"
    ws["L5"] = 2050
    ws["D5"] = date(2026, 5, 5)
    ws["N5"] = date(2026, 3, 25)
    ws["U5"] = date(2026, 4, 16)
    ws["V5"] = date(2026, 4, 20)
    ws["AE5"] = date(2026, 4, 25)
