"""A normal tabular sheet, but the e2e test will pass an empty name_map so
apply_plan can't resolve labels -> canonical fields.

Tests apply_plan's behavior when CanonicalNameMap is missing entries.
"""
from openpyxl import Workbook


def build(wb: Workbook) -> None:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    ws = wb.create_sheet("S")
    ws["A1"] = "IO NO"; ws["B1"] = "STYLE"
    ws["A2"] = "1063"; ws["B2"] = "DWJE"
