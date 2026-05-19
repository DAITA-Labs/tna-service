"""_sample_values returns ≤k non-null samples per identity column."""
import openpyxl
from app.enums.pli_mode import PliMode
from app.enums.row_role import RowRole
from app.models.artifacts import HeaderLabel, RowSpec, SheetPlan
from app.repositories.workbook_repo import register_workbook
from app.services.agents.field_namer import _sample_values


def test_returns_at_most_k_samples_per_label(tmp_path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["B3"] = "IO NO"
    ws["B4"] = 1063
    ws["B5"] = 1064
    ws["B6"] = 1065
    ws["B7"] = 1066
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ctx = register_workbook(p)
    plan = SheetPlan(
        sheet="S", pli_mode=PliMode.ROW_PER_PLI,
        rows=[RowSpec(idx=r, role=RowRole.ANCHOR) for r in (4, 5, 6, 7)],
        header_labels=[HeaderLabel(raw="IO NO", col="B", row=3)],
    )
    samples = _sample_values(ctx.wb["S"], plan, [("IO NO", "B")], k=3)
    assert samples["IO NO"] == [1063, 1064, 1065]


def test_skips_null_cells(tmp_path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["B3"] = "IO NO"
    ws["B4"] = None
    ws["B5"] = 1064
    ws["B6"] = None
    ws["B7"] = 1066
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ctx = register_workbook(p)
    plan = SheetPlan(
        sheet="S", pli_mode=PliMode.ROW_PER_PLI,
        rows=[RowSpec(idx=r, role=RowRole.ANCHOR) for r in (4, 5, 6, 7)],
        header_labels=[HeaderLabel(raw="IO NO", col="B", row=3)],
    )
    samples = _sample_values(ctx.wb["S"], plan, [("IO NO", "B")], k=3)
    assert samples["IO NO"] == [1064, 1066]


def test_returns_empty_for_label_with_no_data_rows(tmp_path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["B3"] = "X"
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ctx = register_workbook(p)
    plan = SheetPlan(sheet="S", pli_mode=PliMode.ROW_PER_PLI, header_labels=[])
    samples = _sample_values(ctx.wb["S"], plan, [("X", "B")], k=3)
    assert samples.get("X") in (None, [])
