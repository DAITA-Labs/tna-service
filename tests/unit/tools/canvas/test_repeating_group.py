"""RepeatingRowGroup detector — section-style header rhythm."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.artifacts.structure import RepeatingRowGroup
from app.tools.canvas.build import build_canvas
from app.tools.canvas.repeating_group import find_repeating_row_groups


DATASET = Path(__file__).resolve().parents[4] / "dataset"


@pytest.fixture
def guess_canvas():
    """GUESS — SECTION_PER_PLI; many repeating header rows."""
    wb = load_workbook(DATASET / "GUESS ATHLEISURE - MAIN FALL 26 MASTER FILE #1.xlsx",
                       data_only=True)
    return build_canvas(wb.active)


@pytest.fixture
def dkn_canvas():
    """DKN — ROW_PER_PLI with no repeating headers."""
    wb = load_workbook(DATASET / "20260129 DKN AW26 DROP 2 WOMEN NOS CK PRO STATUS.xlsx",
                       data_only=True)
    return build_canvas(wb.active)


def test_groups_returned_as_records(guess_canvas) -> None:
    """find_repeating_row_groups returns RepeatingRowGroup records."""
    groups = find_repeating_row_groups(guess_canvas)
    assert isinstance(groups, list)
    assert all(isinstance(g, RepeatingRowGroup) for g in groups)


def test_each_group_has_at_least_two_rows(guess_canvas) -> None:
    """A 'repeating' group by construction has ≥2 member rows."""
    groups = find_repeating_row_groups(guess_canvas)
    for g in groups:
        assert len(g.row_indices) >= 2


def test_guess_has_repeating_header_pattern(guess_canvas) -> None:
    """GUESS's section-style layout produces at least one repeating row group."""
    groups = find_repeating_row_groups(guess_canvas)
    # GUESS has ~15 repeating section headers; expect at least one group with
    # several members
    has_large_group = any(len(g.row_indices) >= 3 for g in groups)
    assert has_large_group


def test_dkn_has_no_or_few_repeating_groups(dkn_canvas) -> None:
    """DKN is ROW_PER_PLI with no section headers — few or zero groups."""
    groups = find_repeating_row_groups(dkn_canvas)
    # ROW_PER_PLI sheets shouldn't trigger many large repeating groups
    large_groups = [g for g in groups if len(g.row_indices) >= 5]
    assert len(large_groups) == 0


def test_row_indices_are_1_indexed(guess_canvas) -> None:
    """row_indices are 1-indexed positions matching the sheet's row numbers."""
    groups = find_repeating_row_groups(guess_canvas)
    for g in groups:
        for r in g.row_indices:
            assert r >= 1


def test_signature_is_string(guess_canvas) -> None:
    """signature is a non-empty string built from the row's cell contents."""
    groups = find_repeating_row_groups(guess_canvas)
    for g in groups:
        assert isinstance(g.signature, str)
        assert g.signature != ""


def test_tool_registered() -> None:
    from app.tools._registry import TOOL_REGISTRY

    assert "find_repeating_row_groups" in TOOL_REGISTRY.names()
