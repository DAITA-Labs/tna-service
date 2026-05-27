"""DateStrip detection — vertical + horizontal runs."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.artifacts.structure import DateStrip
from app.tools.canvas.build import build_canvas
from app.tools.canvas.strips_date import (
    find_date_strips,
    find_date_strips_horizontal,
    find_date_strips_vertical,
)


DATASET = Path(__file__).resolve().parents[4] / "dataset"


@pytest.fixture
def dkn_canvas():
    """DKN baseline — ROW_PER_PLI with vertical stage date columns."""
    wb = load_workbook(DATASET / "20260129 DKN AW26 DROP 2 WOMEN NOS CK PRO STATUS.xlsx",
                       data_only=True)
    return build_canvas(wb.active)


@pytest.fixture
def kv_canvas():
    """63261-TNA — SHEET_IS_PLI with horizontal stage date rows in bands."""
    wb = load_workbook(DATASET / "63261-TNA.xlsx", data_only=True)
    return build_canvas(wb.active)


def test_vertical_date_strips_on_dkn(dkn_canvas) -> None:
    """DKN has multiple vertical date strips (one per stage column)."""
    strips = find_date_strips_vertical(dkn_canvas)
    assert isinstance(strips, list)
    assert all(isinstance(s, DateStrip) for s in strips)
    assert all(s.orientation == "vertical" for s in strips)


def test_horizontal_date_strips_on_kv_layout(kv_canvas) -> None:
    """63261 has horizontal date rows inside its TNA bands (rows 9, 14, 19)."""
    strips = find_date_strips_horizontal(kv_canvas)
    assert all(s.orientation == "horizontal" for s in strips)
    # 63261 has at least one horizontal date strip (≥3 dates in one row)
    assert len(strips) >= 1


def test_combined_returns_both_orientations(kv_canvas) -> None:
    """find_date_strips returns the union of vertical + horizontal."""
    all_strips = find_date_strips(kv_canvas)
    vert = find_date_strips_vertical(kv_canvas)
    horiz = find_date_strips_horizontal(kv_canvas)
    assert len(all_strips) == len(vert) + len(horiz)


def test_min_run_threshold_filters_short_runs(dkn_canvas) -> None:
    """A high min_run threshold should filter out short runs."""
    short_strips = find_date_strips_vertical(dkn_canvas, min_run=3)
    long_strips = find_date_strips_vertical(dkn_canvas, min_run=100)
    assert len(long_strips) <= len(short_strips)


def test_rect_is_1_indexed_inclusive(dkn_canvas) -> None:
    """Each strip's rect uses 1-indexed inclusive coordinates."""
    strips = find_date_strips_vertical(dkn_canvas)
    for s in strips:
        assert s.rect.r0 >= 1
        assert s.rect.c0 >= 1
        assert s.rect.r1 >= s.rect.r0
        # vertical strip: same column
        assert s.rect.c0 == s.rect.c1


def test_tools_registered() -> None:
    from app.tools._registry import TOOL_REGISTRY

    names = TOOL_REGISTRY.names()
    assert "find_date_strips" in names
    assert "find_date_strips_vertical" in names
    assert "find_date_strips_horizontal" in names
