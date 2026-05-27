"""build_canvas tool — channel population + merge handling smoke tests."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.artifacts.canvas import GridCanvas
from app.tools.canvas.build import build_canvas


DATASET = Path(__file__).resolve().parents[4] / "dataset"


@pytest.fixture
def dkn_canvas() -> GridCanvas:
    """Build a canvas from the DKN baseline fixture (small ROW_PER_PLI sheet)."""
    wb = load_workbook(DATASET / "20260129 DKN AW26 DROP 2 WOMEN NOS CK PRO STATUS.xlsx",
                       data_only=True)
    return build_canvas(wb.active)


def test_build_returns_grid_canvas(dkn_canvas: GridCanvas) -> None:
    """build_canvas must return a GridCanvas instance with sheet dimensions."""
    assert isinstance(dkn_canvas, GridCanvas)
    assert dkn_canvas.n_rows > 0
    assert dkn_canvas.n_cols > 0


def test_expected_channels_present(dkn_canvas: GridCanvas) -> None:
    """All ~20 canvas channels must be populated by a single build."""
    expected = {
        "dtype", "density", "plan_marker", "fill_color", "border", "merge",
        "date_like", "text_density_per_row", "dtype_run_row", "dtype_run_col",
        "date_cluster", "density_cluster", "date_flow",
        "bold", "merge_shape",
        "repeating_row_id", "repeating_col_id",
        "empty_row", "empty_col",
        "fill_color_cluster", "column_dtype_dominant",
    }
    assert expected.issubset(set(dkn_canvas.channels.keys()))


def test_merge_ranges_is_set_of_tuples(dkn_canvas: GridCanvas) -> None:
    """merge_ranges is a set of (r0, c0, r1, c1) tuples per canonical artifact."""
    assert isinstance(dkn_canvas.merge_ranges, set)
    for mr in dkn_canvas.merge_ranges:
        assert isinstance(mr, tuple)
        assert len(mr) == 4


def test_cell_values_grid_shape(dkn_canvas: GridCanvas) -> None:
    """cell_values is an n_rows × n_cols list-of-lists."""
    assert len(dkn_canvas.cell_values) == dkn_canvas.n_rows
    for row in dkn_canvas.cell_values:
        assert len(row) == dkn_canvas.n_cols


def test_dtype_matrix_shape(dkn_canvas: GridCanvas) -> None:
    """Every channel matrix is n_rows × n_cols."""
    dtype = dkn_canvas.channels["dtype"]
    assert len(dtype) == dkn_canvas.n_rows
    assert all(len(r) == dkn_canvas.n_cols for r in dtype)


def test_tool_registered() -> None:
    """build_canvas is registered in the TOOL_REGISTRY under its decorated name."""
    from app.tools._registry import TOOL_REGISTRY

    assert "build_canvas" in TOOL_REGISTRY.names()


def test_max_row_caps_dimension() -> None:
    """The optional max_row arg caps the canvas height."""
    wb = load_workbook(DATASET / "20260129 DKN AW26 DROP 2 WOMEN NOS CK PRO STATUS.xlsx",
                       data_only=True)
    canvas = build_canvas(wb.active, max_row=10)
    assert canvas.n_rows == 10
