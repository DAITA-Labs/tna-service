"""Tests for app/services/applier/patterns/*."""
from openpyxl import Workbook
from app.repositories.workbook_repo import register_workbook, clear_cache
from app.models.artifacts import PLIBoundaries
from app.enums.boundary_pattern import BoundaryPattern
# Importing this package triggers handler registration via @pattern_handler.
import app.services.applier.patterns  # noqa: F401
from app.services.applier._registry import get_pattern_handler


def test_one_row_per_pli_iterates_inclusive(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active
    ws["A4"] = 1; ws["A5"] = 2; ws["A6"] = 3
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    b = PLIBoundaries(sheet="Sheet", pattern=BoundaryPattern.ONE_ROW_PER_PLI,
                     data_start_row=4, data_end_row=6, confidence=1.0)
    assert get_pattern_handler("one_row_per_pli")(ctx, b) == [4, 5, 6]


def test_data_then_total_excludes_indicator(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active
    ws["A4"] = "PLI1"; ws["A5"] = "Grand Total"; ws["A6"] = "PLI2"
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    b = PLIBoundaries(sheet="Sheet", pattern=BoundaryPattern.DATA_THEN_TOTAL,
                     data_start_row=4, data_end_row=6,
                     total_row_indicator_col="A",
                     total_row_indicator_value="Grand Total", confidence=1.0)
    assert get_pattern_handler("data_then_total")(ctx, b) == [4, 6]


def test_vertical_merge_iterates_all_rows(tmp_path):
    """vertical_merge iterates every row in [start, end] inclusive — sub-rows
    of merge groups are distinct PLIs (multi-color)."""
    clear_cache()
    wb = Workbook(); ws = wb.active
    ws["B4"] = 1063; ws["B6"] = 1064
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    b = PLIBoundaries(sheet="Sheet", pattern=BoundaryPattern.VERTICAL_MERGE,
                     data_start_row=4, data_end_row=6,
                     grouping_columns=["B"], confidence=1.0)
    assert get_pattern_handler("vertical_merge")(ctx, b) == [4, 5, 6]


def test_one_sheet_per_pli_returns_empty_handler():
    b = PLIBoundaries(sheet="rep", pattern=BoundaryPattern.ONE_SHEET_PER_PLI,
                     sheet_iter=["S1", "S2"], confidence=1.0)
    assert get_pattern_handler("one_sheet_per_pli")(None, b) == []
