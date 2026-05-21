# tests/unit/applier/test_apply_plan_section_per_pli.py
from datetime import datetime
from openpyxl import Workbook
from app.repositories.workbook_repo import register_workbook, clear_cache
from app.models.artifacts import (
    SheetPlan, PliBlock, KVAnchor, StageBandSpec, CanonicalNameMap,
)
from app.enums.pli_mode import PliMode
from app.enums.stage_scope import StageScope
from app.components.applier import apply_plan


def test_section_per_pli_two_blocks(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active; ws.title = "S"
    ws["A4"] = "IO"; ws["B4"] = 1063
    ws["C8"] = "Cut"; ws["D8"] = "Sew"
    ws["C9"] = datetime(2026, 3, 12); ws["D9"] = datetime(2026, 4, 3)
    ws["A14"] = "IO"; ws["B14"] = 1064
    ws["C18"] = "Cut"; ws["D18"] = "Sew"
    ws["C19"] = datetime(2026, 4, 16); ws["D19"] = datetime(2026, 4, 25)
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)

    plan = SheetPlan(
        sheet="S", pli_mode=PliMode.SECTION_PER_PLI,
        pli_blocks=[
            PliBlock(id=0, bbox=(4, 10),
                identity=[KVAnchor(label_cell="A4", value_cell="B4", field="IO")],
                stage_bands=[StageBandSpec(
                    name="ProdBand", name_cell="C8", sub_header_row=8,
                    sub_rows={"plan": 9},
                    stage_cols={"Cut": "C", "Sew": "D"},
                    layout_mode="wide_sub_columns")]),
            PliBlock(id=1, bbox=(14, 20),
                identity=[KVAnchor(label_cell="A14", value_cell="B14", field="IO")],
                stage_bands=[StageBandSpec(
                    name="ProdBand", name_cell="C18", sub_header_row=18,
                    sub_rows={"plan": 19},
                    stage_cols={"Cut": "C", "Sew": "D"},
                    layout_mode="wide_sub_columns")]),
        ],
        stage_scope=StageScope.PLI_LOCAL, confidence=1.0,
    )
    name_map = CanonicalNameMap(
        field_labels={"IO": "io_number"},
        stage_names={"Cut": "cutting", "Sew": "sewing"},
    )
    plis = apply_plan(ctx, plan, name_map)
    assert len(plis) == 2
    assert plis[0].io_number == "1063"
    assert plis[1].io_number == "1064"
    assert any(s.name == "cutting" for s in plis[0].stages)
