"""_coerce must demote non-numeric strings on int canonical fields to None.

Surfaced by make eval crashing on a SHEET_IS_PLI file where a KV anchor
mapped a label to `quantity` but the cell value was the string 'PO'.
"""
import openpyxl

from app.enums.pli_mode import PliMode
from app.models.artifacts import CanonicalNameMap, KVAnchor, SheetPlan
from app.repositories.workbook_repo import register_workbook
from app.components.applier import _coerce, apply_plan


def test_coerce_int_field_with_int_input() -> None:
    assert _coerce("quantity", 1420) == 1420


def test_coerce_int_field_with_float_input() -> None:
    assert _coerce("quantity", 1420.0) == 1420


def test_coerce_int_field_with_numeric_string_input() -> None:
    assert _coerce("quantity", "1,420") == 1420
    assert _coerce("quantity", "  500 ") == 500


def test_coerce_int_field_with_unparseable_string_returns_none() -> None:
    assert _coerce("quantity", "PO") is None
    assert _coerce("order_quantity", "n/a") is None


def test_coerce_int_field_with_bool_returns_none() -> None:
    assert _coerce("quantity", True) is None


def test_apply_plan_routes_bad_quantity_to_metadata(tmp_path, capsys) -> None:
    """End-to-end: SHEET_IS_PLI with a non-numeric quantity cell does not crash."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "QTY"
    ws["B1"] = "PO"  # Not a real quantity — surfaces the bug.
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ctx = register_workbook(p)
    plan = SheetPlan(
        sheet="S", pli_mode=PliMode.SHEET_IS_PLI,
        kv_anchors=[KVAnchor(label_cell="A1", value_cell="B1", field="QTY")],
    )
    nm = CanonicalNameMap(field_labels={"QTY": "quantity"})
    plis = apply_plan(ctx, plan, nm)
    assert len(plis) == 1
    pli = plis[0]
    # quantity must be None (Pydantic accepts None for int | None) — NOT "PO".
    assert pli.quantity is None
    # The original string lands in metadata under the canonical name.
    assert pli.metadata.get("quantity") == "PO"
    # A warning was emitted to stdout via structlog's PrintLoggerFactory.
    out = capsys.readouterr().out
    assert "canonical_coerce_failed" in out
