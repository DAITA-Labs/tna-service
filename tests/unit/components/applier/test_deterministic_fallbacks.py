"""Deterministic fallbacks in apply_plan — _coerce float→int + io_number promotion."""
import openpyxl

from app.enums.pli_mode import PliMode
from app.enums.row_role import RowRole
from app.models.artifacts import (
    CanonicalNameMap, HeaderLabel, KVAnchor, RowSpec, SheetPlan,
)
from app.repositories.workbook_repo import register_workbook
from app.components.per_sheet.applier import (
    _coerce, _promote_io_number_fallback, apply_plan,
)


def test_coerce_io_number_strips_float_zero() -> None:
    """Float 1078.0 → string '1078' (not '1078.0')."""
    assert _coerce("io_number", 1078.0) == "1078"


def test_coerce_io_number_preserves_real_fractional() -> None:
    """Float 1078.5 keeps the fractional portion (real decimal, not Excel artifact)."""
    assert _coerce("io_number", 1078.5) == "1078.5"


def test_coerce_io_number_pass_through_int() -> None:
    """Int 1078 → string '1078'."""
    assert _coerce("io_number", 1078) == "1078"


def test_coerce_io_number_pass_through_string() -> None:
    """String '1078' stays '1078'."""
    assert _coerce("io_number", "1078") == "1078"


def test_coerce_style_code_strips_float_zero() -> None:
    """Same float→int promotion applies to other _STRING_FIELDS."""
    assert _coerce("style_code", 890162.0) == "890162"


def test_promote_io_number_from_metadata_when_io_empty() -> None:
    """When io_number is None and metadata.buyer_po_no exists, promote it."""
    values = {"io_number": None, "metadata": {"buyer_po_no": "PO-00045"}}
    _promote_io_number_fallback(values)
    assert values["io_number"] == "PO-00045"
    assert "buyer_po_no" not in values["metadata"]


def test_promote_no_op_when_io_already_set() -> None:
    """If io_number is already populated, leave metadata alone."""
    values = {"io_number": "1063", "metadata": {"buyer_po_no": "PO-00045"}}
    _promote_io_number_fallback(values)
    assert values["io_number"] == "1063"
    assert values["metadata"]["buyer_po_no"] == "PO-00045"


def test_promote_no_op_when_no_buyer_po_no() -> None:
    """If metadata lacks buyer_po_no, no-op."""
    values = {"io_number": None, "metadata": {"buyer": "ACME"}}
    _promote_io_number_fallback(values)
    assert values.get("io_number") is None


def test_promote_converts_non_string_po_to_string(tmp_path) -> None:
    """If buyer_po_no value is an int, convert to string when promoting."""
    values = {"io_number": None, "metadata": {"buyer_po_no": 131673}}
    _promote_io_number_fallback(values)
    assert values["io_number"] == "131673"


def test_end_to_end_buyer_po_promotion(tmp_path) -> None:
    """Real apply_plan path: a PLI with buyer_po_no in metadata gets io_number promoted."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "Buyer Po No"
    ws["B1"] = "131673"
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ctx = register_workbook(p)
    plan = SheetPlan(
        sheet="S", pli_mode=PliMode.SHEET_IS_PLI,
        kv_anchors=[KVAnchor(label_cell="A1", value_cell="B1", field="Buyer Po No")],
    )
    # Simulate FieldNamer mapping to buyer_po_no (metadata) — the fallback should
    # promote it to io_number.
    nm = CanonicalNameMap(field_labels={"Buyer Po No": "buyer_po_no"})
    plis = apply_plan(ctx, plan, nm)
    assert plis[0].io_number == "131673"
    assert "buyer_po_no" not in (plis[0].metadata or {})


def test_end_to_end_float_io_number(tmp_path) -> None:
    """A cell holding 1078 (int-valued float from Excel) becomes io_number='1078' not '1078.0'."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "IO"
    ws["B1"] = 1078.0  # float in the Excel cell
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ctx = register_workbook(p)
    plan = SheetPlan(
        sheet="S", pli_mode=PliMode.SHEET_IS_PLI,
        kv_anchors=[KVAnchor(label_cell="A1", value_cell="B1", field="IO")],
    )
    nm = CanonicalNameMap(field_labels={"IO": "io_number"})
    plis = apply_plan(ctx, plan, nm)
    assert plis[0].io_number == "1078"
