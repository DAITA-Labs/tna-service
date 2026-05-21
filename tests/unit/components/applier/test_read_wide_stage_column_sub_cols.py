"""_read_wide_stage_column iterates sub_columns and routes them correctly."""
from datetime import date
import openpyxl
from app.enums.pli_mode import PliMode
from app.enums.row_role import RowRole
from app.models.artifacts import (
    CanonicalNameMap, HeaderLabel, RowSpec, SheetPlan, StageBandSpec, StageColumn,
)
from app.repositories.workbook_repo import register_workbook
from app.components.applier import apply_plan


def test_sub_column_routes_to_stage_metadata(tmp_path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["U3"] = "CUTTING"
    ws["V3"] = "Actual"
    ws["U4"] = date(2026, 3, 12)
    ws["V4"] = date(2026, 3, 16)
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ctx = register_workbook(p)
    band = StageBandSpec(
        name="b", name_cell="U3", sub_header_row=3,
        sub_rows={"plan": 4},
        layout_mode="wide_sub_columns",
        stage_columns=[
            StageColumn(name="CUTTING", name_cell="U3", primary_col="U",
                         sub_columns={"Actual": "V"}),
        ],
    )
    plan = SheetPlan(
        sheet="S", pli_mode=PliMode.ROW_PER_PLI,
        header_rows=[3], stage_bands=[band],
        rows=[RowSpec(idx=4, role=RowRole.ANCHOR)],
        header_labels=[HeaderLabel(raw="x", col="A", row=3)],  # to satisfy invariant
    )
    name_map = CanonicalNameMap(
        stage_names={"CUTTING": "cutting"},
        stage_subfield_labels={"Actual": "actual_date"},
    )
    plis = apply_plan(ctx, plan, name_map)
    assert len(plis) == 1
    assert len(plis[0].stages) == 1
    stage = plis[0].stages[0]
    assert stage.name == "cutting"
    assert stage.planned_date == date(2026, 3, 12)
    assert stage.metadata.get("actual_date") == date(2026, 3, 16)


def test_ignore_sentinel_skips_sub_column(tmp_path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["U3"] = "CUTTING"
    ws["V3"] = "Remarks"
    ws["U4"] = date(2026, 3, 12)
    ws["V4"] = "n/a"
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ctx = register_workbook(p)
    band = StageBandSpec(
        name="b", name_cell="U3", sub_header_row=3,
        sub_rows={"plan": 4},
        layout_mode="wide_sub_columns",
        stage_columns=[
            StageColumn(name="CUTTING", name_cell="U3", primary_col="U",
                         sub_columns={"Remarks": "V"}),
        ],
    )
    plan = SheetPlan(
        sheet="S", pli_mode=PliMode.ROW_PER_PLI,
        header_rows=[3], stage_bands=[band],
        rows=[RowSpec(idx=4, role=RowRole.ANCHOR)],
        header_labels=[HeaderLabel(raw="x", col="A", row=3)],
    )
    name_map = CanonicalNameMap(
        stage_names={"CUTTING": "cutting"},
        stage_subfield_labels={"Remarks": "ignore"},
    )
    plis = apply_plan(ctx, plan, name_map)
    stage = plis[0].stages[0]
    assert "Remarks" not in stage.metadata
    assert "remarks" not in stage.metadata
