"""_collect_header_labels — populates header_labels for ROW_PER_PLI, skips claimed cols."""
import openpyxl
from app.enums.pli_mode import PliMode
from app.models.artifacts import HeaderLabel, SheetPlan, StageBandSpec, StageColumn
from app.components.planner.plan import _collect_header_labels


def _ws_with_header(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    # Row 3 header row
    ws["B3"] = "IO NO"
    ws["F3"] = "STYLE"
    ws["K3"] = "COLOR"
    ws["U3"] = "CUTTING"   # claimed by stage band
    ws["V3"] = "Actual"    # claimed sub-col
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    return openpyxl.load_workbook(p)["S"]


def test_collects_identity_columns_skipping_stage_cols(tmp_path) -> None:
    ws = _ws_with_header(tmp_path)
    band = StageBandSpec(
        name="b", name_cell="U3", sub_header_row=3,
        stage_columns=[StageColumn(name="CUTTING", name_cell="U3", primary_col="U",
                                    sub_columns={"Actual": "V"})],
    )
    plan = SheetPlan(sheet="S", pli_mode=PliMode.ROW_PER_PLI,
                     header_rows=[3], stage_bands=[band])
    labels = _collect_header_labels(ws, plan)
    raws = sorted(hl.raw for hl in labels)
    assert raws == ["COLOR", "IO NO", "STYLE"]
    assert all(isinstance(hl, HeaderLabel) for hl in labels)
    by_raw = {hl.raw: hl for hl in labels}
    assert by_raw["IO NO"].col == "B"
    assert by_raw["IO NO"].row == 3


def test_returns_empty_for_non_row_per_pli(tmp_path) -> None:
    ws = _ws_with_header(tmp_path)
    plan = SheetPlan(sheet="S", pli_mode=PliMode.SHEET_IS_PLI, header_rows=[3])
    assert _collect_header_labels(ws, plan) == []
    plan2 = SheetPlan(sheet="S", pli_mode=PliMode.SECTION_PER_PLI, header_rows=[3])
    assert _collect_header_labels(ws, plan2) == []


def test_collects_from_row_below_when_title_row_dominates(tmp_path) -> None:
    """When SheetRowPlanner detects a title-only header row, labels come from the next row.

    Mirrors the DKN / Northern Reflections pattern: row 1 is a wide merged title in
    column A; real column labels live in row 2; data starts in row 3. The planner must
    return the real labels from row 2, not just the title string.

    The title-row correction is driven by _detect_title_row_extra_header +
    _apply_extra_header_row in SheetRowPlanner.run(). This test calls those helpers
    directly and then verifies _collect_header_labels with the corrected plan.
    """
    import openpyxl
    from app.enums.pli_mode import PliMode
    from app.enums.row_role import RowRole
    from app.models.artifacts import RowSpec, SheetPlan
    from app.components.planner.plan import (
        _apply_extra_header_row,
        _collect_header_labels,
        _detect_title_row_extra_header,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    # Row 1: a long title in col A only (merged across A–H).
    ws["A1"] = "DKN AW26 DROP 2 WOMEN NOSORDER PRODUCTION STATUS REPORT"
    ws.merge_cells("A1:H1")
    # Row 2: actual column labels.
    ws["A2"] = "S.No"
    ws["B2"] = "IO No"
    ws["C2"] = "Style No"
    ws["D2"] = "Color"
    ws["E2"] = "Quantity"
    # Row 3: data rows.
    ws["A3"] = 1
    ws["B3"] = "1063"
    ws["C3"] = "STYLE-A"
    ws["D3"] = "Red"
    ws["E3"] = 100
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ws = openpyxl.load_workbook(p)["S"]

    # Simulate the planner having classified row 1 as HEADER (title-row
    # misidentification) and row 2 as ANCHOR (the real header row).
    rows = [
        RowSpec(idx=1, role=RowRole.HEADER),
        RowSpec(idx=2, role=RowRole.ANCHOR, group_id=0),
        RowSpec(idx=3, role=RowRole.ANCHOR, group_id=1),
    ]
    plan = SheetPlan(sheet="S", pli_mode=PliMode.ROW_PER_PLI, header_rows=[1], rows=rows)

    # The title-row detector should return row 2 as the extra header row.
    extra_row = _detect_title_row_extra_header(ws, plan)
    assert extra_row == 2, f"Expected extra_row=2 but got {extra_row}"

    # Apply the correction and pass only the true header row to label collector.
    corrected_rows = _apply_extra_header_row(plan.rows, extra_row)
    new_header_idxs = {r.idx for r in corrected_rows if r.role is RowRole.HEADER}
    plan = plan.model_copy(update={
        "rows": corrected_rows,
        "header_rows": sorted(new_header_idxs),
    })
    labels_plan = plan.model_copy(update={"header_rows": [extra_row]})
    labels = _collect_header_labels(ws, labels_plan)
    raws = {hl.raw for hl in labels}
    # Must contain the real labels from row 2, not just the title string from row 1.
    assert {"IO No", "Style No", "Color", "Quantity"}.issubset(raws), (
        f"Expected real column labels but got: {raws}"
    )
