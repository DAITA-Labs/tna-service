"""Stage-band detector — wide_sub_columns shape via stage_columns list."""
from datetime import date

from app.models.artifacts import SheetSignals
from app.repositories.workbook_repo import register_workbook
from app.components.planner.stage_band_detector import detect_stage_bands


def _signals(max_row: int, max_col: int) -> SheetSignals:
    return SheetSignals(sheet="S", max_row=max_row, max_col=max_col)


def test_wide_band_emits_stage_columns_list(tmp_path) -> None:
    """Detector populates stage_columns mirroring stage_cols (empty sub_columns)."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A3"] = "S NO"
    ws["B3"] = "IO"
    ws["U3"] = "CUTTING"
    ws["V3"] = "SEWING"
    ws["U4"] = date(2026, 3, 12)
    ws["V4"] = date(2026, 4, 3)
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ctx = register_workbook(p)

    bands = detect_stage_bands(ctx, "S", _signals(max_row=4, max_col=22))
    assert len(bands) == 1
    b = bands[0]
    assert b.layout_mode == "wide_sub_columns"
    assert b.stage_cols == {"CUTTING": "U", "SEWING": "V"}
    assert len(b.stage_columns) == 2
    by_name = {sc.name: sc for sc in b.stage_columns}
    assert by_name["CUTTING"].primary_col == "U"
    assert by_name["CUTTING"].name_cell == "U3"
    assert by_name["CUTTING"].sub_columns == {}
    assert by_name["SEWING"].primary_col == "V"


def test_wide_band_fills_sub_columns_from_row_below(tmp_path) -> None:
    """When row below stage name has labels, they populate sub_columns."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    # Row 3 = stage names, Row 4 = sub-field names, Row 5 = data row with dates
    ws["U3"] = "CUTTING"
    ws["V3"] = None  # no second-stage-name; V belongs under CUTTING
    ws["U4"] = "Plan"
    ws["V4"] = "Actual"
    ws["U5"] = date(2026, 3, 12)
    ws["V5"] = date(2026, 3, 16)
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ctx = register_workbook(p)

    bands = detect_stage_bands(ctx, "S", _signals(max_row=5, max_col=22))
    # The detector seeds bands from rows with ≥2 date cells; row 5 qualifies.
    # The header row above (row 4) holds the immediate sub-headers; row 3 holds
    # the stage name "CUTTING" with V3 empty (shared stage span).
    assert any(b.layout_mode == "wide_sub_columns" for b in bands)
    band = next(b for b in bands if b.layout_mode == "wide_sub_columns")
    by_name = {sc.name: sc for sc in band.stage_columns}
    assert "CUTTING" in by_name
    # Sub-column "Actual" should be on CUTTING with V col.
    assert by_name["CUTTING"].sub_columns.get("Actual") == "V"
