"""populate_patterns / populate_semantics — bag wiring helpers."""
from __future__ import annotations

import openpyxl

from app.artifacts.structure import StructureBag
from app.components.structure.populate import populate_patterns, populate_semantics
from app.tools.canvas.build import build_canvas


def _tabular_canvas():
    """Make a small tabular workbook: header row + 5 PLI rows + a date column."""
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["IO No", "Style", "Qty", "Delivery"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = openpyxl.styles.Font(bold=True)
    import datetime as dt

    for r in range(3, 8):
        ws.cell(row=r, column=1, value=f"IO-{r}")
        ws.cell(row=r, column=2, value=f"S{r}")
        ws.cell(row=r, column=3, value=100 + r)
        ws.cell(row=r, column=4, value=dt.date(2026, 5, r))
    return build_canvas(ws)


def test_populate_patterns_fills_all_bag_lists() -> None:
    """populate_patterns runs every detector — bag's pattern lists are populated where signals exist."""
    canvas = _tabular_canvas()
    bag = StructureBag()
    populate_patterns(canvas, bag)
    # Date column → at least one DateStrip
    assert len(bag.date_strips) >= 1
    # Qty column → at least one IntStrip
    assert len(bag.int_strips) >= 1
    # Bold header row → at least one BoldStrip
    assert len(bag.bold_strips) >= 1


def test_populate_patterns_is_additive_does_not_reset() -> None:
    """populate_patterns appends — calling twice doubles records (caller controls fresh bag)."""
    canvas = _tabular_canvas()
    bag = StructureBag()
    populate_patterns(canvas, bag)
    n_dates_first = len(bag.date_strips)
    populate_patterns(canvas, bag)
    assert len(bag.date_strips) == 2 * n_dates_first


def test_populate_semantics_runs_after_patterns() -> None:
    """populate_semantics fills bag.header_band and bag.data_row_ranges on a tabular sheet."""
    canvas = _tabular_canvas()
    bag = StructureBag()
    populate_patterns(canvas, bag)
    populate_semantics(canvas, bag)
    # Header on row 2 with bold "IO No" / "Style" should anchor a band
    assert bag.header_band is not None
    # Data rows 3..7 should produce at least one DataRowRange
    assert len(bag.data_row_ranges) >= 1
    rng = bag.data_row_ranges[0]
    assert rng.row_start >= 3
    assert rng.row_end >= rng.row_start


def test_populate_semantics_idempotent_when_resolvers_idempotent() -> None:
    """Running populate_semantics twice with the same patterns shouldn't crash (resolvers may overwrite)."""
    canvas = _tabular_canvas()
    bag = StructureBag()
    populate_patterns(canvas, bag)
    populate_semantics(canvas, bag)
    band_first = bag.header_band
    populate_semantics(canvas, bag)
    # Header band still set; resolvers themselves decide overwrite vs append semantics
    assert bag.header_band is not None
    assert bag.header_band == band_first
