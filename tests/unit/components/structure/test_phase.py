"""StructurePhase orchestrator — end-to-end sheet → LayoutHint smoke."""
from __future__ import annotations

import datetime as dt

import openpyxl
from openpyxl.styles import Font

from app.artifacts.canvas import GridCanvas
from app.artifacts.layout import LayoutHint
from app.artifacts.structure import StructureBag
from app.components.structure.phase import run_structure_phase


def _tabular_sheet():
    """Header on row 2 (bold) + 5 PLI rows + a date column on col 4."""
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["IO No", "Style", "Qty", "Delivery"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = Font(bold=True)
    for r in range(3, 8):
        ws.cell(row=r, column=1, value=f"IO-{r}")
        ws.cell(row=r, column=2, value=f"S{r}")
        ws.cell(row=r, column=3, value=100 + r)
        ws.cell(row=r, column=4, value=dt.date(2026, 5, r))
    return ws


def test_run_structure_phase_returns_triple() -> None:
    """The orchestrator returns (canvas, bag, hint) with correct types."""
    canvas, bag, hint = run_structure_phase(_tabular_sheet())
    assert isinstance(canvas, GridCanvas)
    assert isinstance(bag, StructureBag)
    assert isinstance(hint, LayoutHint)


def test_run_structure_phase_populates_bag_end_to_end() -> None:
    """Bag carries both pattern records (date strips) and semantic records (header band)."""
    _, bag, _ = run_structure_phase(_tabular_sheet())
    # Pattern records populated
    assert len(bag.date_strips) >= 1
    assert len(bag.int_strips) >= 1
    # Semantic records populated
    assert bag.header_band is not None
    assert len(bag.data_row_ranges) >= 1


def test_run_structure_phase_hint_carries_axes_and_cluster_id() -> None:
    """LayoutHint carries the cluster_id and a populated LayoutAxes."""
    _, _, hint = run_structure_phase(_tabular_sheet(), cluster_id="cluster_main")
    assert hint.cluster_id == "cluster_main"
    assert hint.axes.pli_axis in ("vertical", "horizontal", "sheet", "sectional")
    assert set(hint.axes.confidence.keys()) == {"pli", "stage", "subfield"}


def test_run_structure_phase_hint_pre_narrows_candidates_for_tabular() -> None:
    """Tabular sheet yields shared candidate_rows and identifier candidate_columns."""
    _, _, hint = run_structure_phase(_tabular_sheet())
    # candidate_rows should cover data rows for every identifier canonical
    assert "io_number" in hint.candidate_rows
    rows = hint.candidate_rows["io_number"]
    assert any(r in rows for r in range(3, 8))
    # candidate_columns should have at least one identifier canonical mapped to col indices
    assert len(hint.candidate_columns) >= 1


def test_run_structure_phase_default_cluster_id() -> None:
    """When no cluster_id is passed, it defaults to 'default'."""
    _, _, hint = run_structure_phase(_tabular_sheet())
    assert hint.cluster_id == "default"
