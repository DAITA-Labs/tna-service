"""StructurePhase Haystack component wraps run_structure_phase."""
from __future__ import annotations

import datetime as dt

import openpyxl
from haystack import Pipeline
from openpyxl.styles import Font

from app.artifacts.canvas import GridCanvas
from app.artifacts.layout import LayoutHint
from app.artifacts.structure import StructureBag
from app.components.structure.phase import StructurePhase


def _tabular_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in enumerate(["IO No", "Style", "Qty", "Delivery"], start=1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = Font(bold=True)
    for r in range(3, 8):
        ws.cell(row=r, column=1, value=f"IO-{r}")
        ws.cell(row=r, column=2, value=f"S{r}")
        ws.cell(row=r, column=3, value=100 + r)
        ws.cell(row=r, column=4, value=dt.date(2026, 5, r))
    return ws


def test_component_run_returns_three_named_outputs() -> None:
    """`run` returns a dict with canvas / bag / hint keys."""
    comp = StructurePhase()
    out = comp.run(sheet=_tabular_sheet(), cluster_id="c0")
    assert set(out.keys()) == {"canvas", "bag", "hint"}
    assert isinstance(out["canvas"], GridCanvas)
    assert isinstance(out["bag"], StructureBag)
    assert isinstance(out["hint"], LayoutHint)
    assert out["hint"].cluster_id == "c0"


def test_component_default_cluster_id() -> None:
    """When cluster_id is omitted, the hint carries 'default'."""
    out = StructurePhase().run(sheet=_tabular_sheet())
    assert out["hint"].cluster_id == "default"


def test_component_exposes_haystack_input_output_sockets() -> None:
    """The @component decorator must register input + output sockets."""
    comp = StructurePhase()
    # Haystack stores the input/output socket metadata on the instance after
    # the decorator runs. We don't invoke the pipeline (its breakpoint
    # snapshot serializer recursion-loops on openpyxl Worksheet inputs); we
    # check the socket metadata directly.
    in_sockets = comp.__haystack_input__._sockets_dict
    out_sockets = comp.__haystack_output__._sockets_dict
    assert "sheet" in in_sockets
    assert "cluster_id" in in_sockets
    assert {"canvas", "bag", "hint"}.issubset(out_sockets.keys())


def test_component_can_be_added_to_a_pipeline() -> None:
    """Pipeline.add_component accepts StructurePhase without raising."""
    pipeline = Pipeline()
    pipeline.add_component("structure", StructurePhase())
    assert "structure" in pipeline.graph.nodes
