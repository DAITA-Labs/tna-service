"""WorkbookPhase orchestrator — workbook → list[ClusterAnchorBundle]."""
from __future__ import annotations

import datetime as dt

import openpyxl
from haystack import Pipeline
from openpyxl.styles import Font

from app.artifacts.workbook import ClusterAnchorBundle
from app.components.workbook.workbook_phase import WorkbookPhase, run_workbook_phase


def _tabular_pli_sheet(ws, prefix: str) -> None:
    """Populate `ws` with a small PLI-shaped table (header + 5 rows + date col)."""
    for c, h in enumerate(["IO No", "Style", "Qty", "Delivery"], start=1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = Font(bold=True)
    for r in range(3, 8):
        ws.cell(row=r, column=1, value=f"{prefix}-{r}")
        ws.cell(row=r, column=2, value=f"S{r}")
        ws.cell(row=r, column=3, value=100 + r)
        ws.cell(row=r, column=4, value=dt.date(2026, 5, r))


def _other_sheet(ws) -> None:
    """A metadata-shaped sheet — no dates, no PLI signals."""
    ws.cell(row=1, column=1, value="Buyer")
    ws.cell(row=1, column=2, value="Acme")
    ws.cell(row=2, column=1, value="Season")
    ws.cell(row=2, column=2, value="FW26")


def _workbook(*sheet_specs):
    """Build a workbook with one sheet per (title, populate_fn) pair."""
    wb = openpyxl.Workbook()
    default = wb.active
    if sheet_specs:
        default.title = sheet_specs[0][0]
        sheet_specs[0][1](default)
    for title, populate in sheet_specs[1:]:
        ws = wb.create_sheet(title=title)
        populate(ws)
    return wb


def test_run_workbook_phase_emits_one_bundle_per_cluster() -> None:
    """Every cluster the clusterer produces becomes a bundle.

    No role filter — operators curate the workbook upstream, so the
    phase passes every cluster through to the planner.
    """
    wb = _workbook(
        ("PLI-1", lambda ws: _tabular_pli_sheet(ws, "P1")),
        ("PLI-2", lambda ws: _tabular_pli_sheet(ws, "P2")),
        ("Info", _other_sheet),
    )
    bundles = run_workbook_phase(wb)

    assert all(isinstance(b, ClusterAnchorBundle) for b in bundles)
    # All clusters surface — including the metadata-shaped one.
    anchors = {b.anchor_sheet_name for b in bundles}
    assert anchors.issuperset({"PLI-1", "Info"}) or anchors.issuperset({"PLI-2", "Info"})


def test_sibling_canvases_populated_for_multi_sheet_clusters() -> None:
    """A cluster grouping multiple sibling sheets carries their canvases on the bundle."""
    wb = _workbook(
        ("PLI-1", lambda ws: _tabular_pli_sheet(ws, "P1")),
        ("PLI-2", lambda ws: _tabular_pli_sheet(ws, "P2")),
    )
    bundles = run_workbook_phase(wb)
    # Both PLI sheets share a signature — they should cluster together.
    multi = [b for b in bundles if len(b.cluster.sheet_names) >= 2]
    assert multi, "expected the two identically-shaped sheets to share a cluster"
    bundle = multi[0]
    expected_siblings = set(bundle.cluster.sheet_names) - {bundle.anchor_sheet_name}
    assert set(bundle.sibling_canvases.keys()) == expected_siblings


def test_bundle_propagates_cluster_id_onto_hint() -> None:
    wb = _workbook(("PLI-A", lambda ws: _tabular_pli_sheet(ws, "PA")))
    bundles = run_workbook_phase(wb)
    assert len(bundles) == 1
    assert bundles[0].hint.cluster_id == bundles[0].cluster.cluster_id


def test_workbook_phase_component_returns_bundles() -> None:
    wb = _workbook(("PLI-A", lambda ws: _tabular_pli_sheet(ws, "PA")))
    out = WorkbookPhase().run(workbook=wb)
    assert set(out.keys()) == {"bundles"}
    assert len(out["bundles"]) == 1


def test_workbook_phase_component_sockets() -> None:
    comp = WorkbookPhase()
    assert "workbook" in comp.__haystack_input__._sockets_dict
    assert "bundles" in comp.__haystack_output__._sockets_dict


def test_workbook_phase_addable_to_pipeline() -> None:
    pipeline = Pipeline()
    pipeline.add_component("workbook_phase", WorkbookPhase())
    assert "workbook_phase" in pipeline.graph.nodes
