"""Haystack @component wrappers: WorkbookProfiler / SheetClusterer."""
from __future__ import annotations

import openpyxl
from haystack import Pipeline

from app.artifacts.workbook import PliCluster, SheetSignature
from app.components.workbook.clusterer import SheetClusterer
from app.components.workbook.profiler import WorkbookProfiler


def _wb_with_sheets(*sheet_specs):
    """Build an openpyxl workbook with one sheet per (title, cells) pair."""
    wb = openpyxl.Workbook()
    default = wb.active
    if sheet_specs:
        default.title = sheet_specs[0][0]
        for r, c, v in sheet_specs[0][1]:
            default.cell(row=r, column=c, value=v)
    for title, cells in sheet_specs[1:]:
        ws = wb.create_sheet(title=title)
        for r, c, v in cells:
            ws.cell(row=r, column=c, value=v)
    return wb


# ─── WorkbookProfiler ───────────────────────────────────────────────────────


def test_workbook_profiler_emits_one_signature_per_sheet() -> None:
    wb = _wb_with_sheets(
        ("PLI-A", [(1, 1, "IO No"), (2, 1, "IO-1")]),
        ("PLI-B", [(1, 1, "Style"), (2, 1, "S001")]),
    )
    out = WorkbookProfiler().run(workbook=wb)
    assert set(out.keys()) == {"signatures"}
    assert len(out["signatures"]) == 2
    assert {s.sheet_name for s in out["signatures"]} == {"PLI-A", "PLI-B"}


def test_workbook_profiler_socket_metadata_registered() -> None:
    comp = WorkbookProfiler()
    assert "workbook" in comp.__haystack_input__._sockets_dict
    assert "signatures" in comp.__haystack_output__._sockets_dict


# ─── SheetClusterer ─────────────────────────────────────────────────────────


def _signature(name: str, mask_cells=((1, 1),)) -> SheetSignature:
    return SheetSignature(
        sheet_name=name, n_rows=2, n_cols=2,
        non_blank_mask=frozenset(mask_cells),
    )


def test_sheet_clusterer_returns_clusters() -> None:
    sigs = [_signature("A"), _signature("B")]
    out = SheetClusterer().run(signatures=sigs)
    assert set(out.keys()) == {"clusters"}
    assert all(isinstance(c, PliCluster) for c in out["clusters"])


def test_sheet_clusterer_threshold_override_constructor_arg() -> None:
    """A SheetClusterer constructed with a low threshold merges aggressively."""
    sigs = [_signature("A"), _signature("B", mask_cells=((2, 2),))]
    strict = SheetClusterer().run(signatures=sigs)["clusters"]
    lenient = SheetClusterer(threshold=0.0).run(signatures=sigs)["clusters"]
    assert len(strict) == 2
    assert len(lenient) == 1


def test_sheet_clusterer_pipeline_addable() -> None:
    pipeline = Pipeline()
    pipeline.add_component("clusterer", SheetClusterer())
    assert "clusterer" in pipeline.graph.nodes
