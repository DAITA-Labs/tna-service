from openpyxl import Workbook
from app.repositories.workbook_repo import register_workbook, clear_cache
from app.models.artifacts import SheetPlan, RowSpec, CanonicalNameMap
from app.enums.pli_mode import PliMode
from app.enums.row_role import RowRole
from app.enums.stage_scope import StageScope
from app.services.applier.apply_plan import apply_plan


def test_row_per_pli_two_anchors_with_children(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active; ws.title = "S"
    ws["A1"] = "IO NO"; ws["B1"] = "COLOR"; ws["C1"] = "QTY"
    ws["A2"] = "1063"; ws["B2"] = "MAGENTA"; ws["C2"] = 2356
    ws["B3"] = "NAVY"; ws["C3"] = 2356
    ws["A4"] = "1064"; ws["B4"] = "PINE"; ws["C4"] = 2050
    ws.merge_cells("A2:A3")
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)

    plan = SheetPlan(
        sheet="S", pli_mode=PliMode.ROW_PER_PLI, identity_column="A",
        header_rows=[1],
        rows=[
            RowSpec(idx=2, role=RowRole.ANCHOR, group_id=0),
            RowSpec(idx=3, role=RowRole.CHILD, anchor_idx=2, group_id=0),
            RowSpec(idx=4, role=RowRole.ANCHOR, group_id=1),
        ],
        stage_scope=StageScope.SHEET_LEVEL, confidence=1.0,
    )
    name_map = CanonicalNameMap(
        field_labels={"IO NO": "io_number", "COLOR": "color_code", "QTY": "quantity"},
    )
    plis = apply_plan(ctx, plan, name_map)
    assert len(plis) == 3
    assert plis[0].io_number == "1063"
    assert plis[0].color_code == "MAGENTA"
    assert plis[1].io_number == "1063"
    assert plis[1].color_code == "NAVY"
    assert plis[2].io_number == "1064"
