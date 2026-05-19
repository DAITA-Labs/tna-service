"""_emit_single_row_pli reads from plan.header_labels (not re-scanning header_rows)."""
import openpyxl
from app.enums.pli_mode import PliMode
from app.enums.row_role import RowRole
from app.models.artifacts import CanonicalNameMap, HeaderLabel, RowSpec, SheetPlan
from app.repositories.workbook_repo import register_workbook
from app.services.applier.apply_plan import apply_plan


def test_canonical_field_via_header_label(tmp_path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["B3"] = "IO NO"
    ws["B4"] = "1063"
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ctx = register_workbook(p)
    plan = SheetPlan(
        sheet="S", pli_mode=PliMode.ROW_PER_PLI,
        header_rows=[3],
        rows=[RowSpec(idx=4, role=RowRole.ANCHOR)],
        header_labels=[HeaderLabel(raw="IO NO", col="B", row=3)],
    )
    name_map = CanonicalNameMap(field_labels={"IO NO": "io_number"})
    plis = apply_plan(ctx, plan, name_map)
    assert len(plis) == 1
    assert plis[0].io_number == "1063"
    # Confidence populated
    assert plis[0].confidence.get("io_number") == 0.85


def test_unknown_label_routes_to_metadata(tmp_path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["C3"] = "BUYER PO"
    ws["C4"] = "PO-00045"
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ctx = register_workbook(p)
    plan = SheetPlan(
        sheet="S", pli_mode=PliMode.ROW_PER_PLI,
        header_rows=[3],
        rows=[RowSpec(idx=4, role=RowRole.ANCHOR)],
        header_labels=[HeaderLabel(raw="BUYER PO", col="C", row=3)],
    )
    name_map = CanonicalNameMap()  # empty → fall back to raw label
    plis = apply_plan(ctx, plan, name_map)
    assert plis[0].metadata.get("BUYER PO") == "PO-00045"
