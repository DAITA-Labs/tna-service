from datetime import datetime
from openpyxl import Workbook
from app.repositories.workbook_repo import register_workbook, clear_cache
from app.models.artifacts import SheetPlan, KVAnchor, StageBandSpec, CanonicalNameMap
from app.enums.pli_mode import PliMode
from app.enums.stage_scope import StageScope
from app.services.applier.apply_plan import apply_plan


def test_sheet_is_pli_one_pli_per_sheet(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active; ws.title = "63315"
    ws["A4"] = "Job No"; ws["B4"] = 63315
    ws["A5"] = "Quantity"; ws["B5"] = 254886
    ws["A8"] = "Pre-Prod TNA"
    ws["C8"] = "L/D send"; ws["D8"] = "Fit send"
    ws["C9"] = datetime(2026, 3, 1)
    ws["D9"] = datetime(2026, 3, 5)
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)

    plan = SheetPlan(
        sheet="63315", pli_mode=PliMode.SHEET_IS_PLI,
        kv_anchors=[
            KVAnchor(label_cell="A4", value_cell="B4", field="Job No"),
            KVAnchor(label_cell="A5", value_cell="B5", field="Quantity"),
        ],
        stage_bands=[StageBandSpec(
            name="Pre-Prod TNA", name_cell="A8", sub_header_row=8,
            sub_rows={"plan": 9},
            stage_cols={"L/D send": "C", "Fit send": "D"},
            layout_mode="wide_sub_columns",
        )],
        stage_scope=StageScope.SHEET_LEVEL, confidence=1.0,
    )
    name_map = CanonicalNameMap(
        field_labels={"Job No": "io_number", "Quantity": "quantity"},
        stage_names={"L/D send": "lab_dip_send", "Fit send": "fit_send"},
    )
    plis = apply_plan(ctx, plan, name_map)
    assert len(plis) == 1
    assert plis[0].io_number == "63315"
    assert plis[0].quantity == 254886
    names = [s.name for s in plis[0].stages]
    assert "lab_dip_send" in names
    assert "fit_send" in names
