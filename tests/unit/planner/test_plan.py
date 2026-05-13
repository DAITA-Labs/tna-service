from datetime import datetime
from openpyxl import Workbook
from app.repositories.workbook_repo import register_workbook, clear_cache
from app.services.planner.plan import SheetRowPlanner
from app.enums.pli_mode import PliMode
from app.enums.row_role import RowRole


def _save(tmp_path, fn):
    wb = Workbook(); ws = wb.active; ws.title = "S"
    fn(ws)
    p = tmp_path / "x.xlsx"; wb.save(p)
    return p


def test_plan_row_per_pli_tabular(tmp_path):
    clear_cache()
    def fill(ws):
        ws["A1"] = "IO NO"; ws["B1"] = "STYLE"
        ws["A2"] = 1063; ws["B2"] = "DWJE"
        ws["A3"] = 1064; ws["B3"] = "DWJF"
    p = _save(tmp_path, fill)
    ctx = register_workbook(p)
    plan = SheetRowPlanner().run(workbook_ctx=ctx, sheet="S")["plan"]
    assert plan.pli_mode is PliMode.ROW_PER_PLI
    assert plan.identity_column == "A"
    anchors = [r for r in plan.rows if r.role is RowRole.ANCHOR]
    assert len(anchors) == 2


def test_plan_sheet_is_pli(tmp_path):
    clear_cache()
    def fill(ws):
        ws["A3"] = "Job No"; ws["B3"] = 63315
        ws["A4"] = "Quantity"; ws["B4"] = 254886
        ws["A8"] = "Pre-Prod TNA"
        ws["C8"] = "L/D send"; ws["D8"] = "Fit send"
        ws["B9"] = "Plan"
        ws["C9"] = datetime(2026, 3, 1); ws["D9"] = datetime(2026, 3, 5)
    p = _save(tmp_path, fill)
    ctx = register_workbook(p)
    plan = SheetRowPlanner().run(workbook_ctx=ctx, sheet="S")["plan"]
    assert plan.pli_mode is PliMode.SHEET_IS_PLI
    assert len(plan.kv_anchors) >= 2
    assert len(plan.stage_bands) >= 1


def test_plan_christian_berg_like(tmp_path):
    """7 PLIs: 2 anchors + 5 children in 2 groups separated by totals."""
    clear_cache()
    def fill(ws):
        ws["A1"] = "S NO"; ws["B1"] = "IO NO"; ws["K1"] = "COLOR"; ws["L1"] = "ORDER QTY"
        ws["A2"] = 1; ws["B2"] = 1063; ws["K2"] = "MAGENTA"; ws["L2"] = 2356
        ws["K3"] = "NAVY"; ws["L3"] = 2356
        ws["K4"] = "WHITE"; ws["L4"] = 2356
        ws["K5"] = "DEEPTAUPE"; ws["L5"] = 2356
        ws["L6"] = 9424
        ws["A7"] = 2; ws["B7"] = 1064; ws["K7"] = "NAVY-OW"; ws["L7"] = 2050
        ws["K8"] = "SMOKE"; ws["L8"] = 1576
        ws["K9"] = "GRAY"; ws["L9"] = 1576
        ws["L10"] = 5202
        ws.merge_cells("A2:A5"); ws.merge_cells("B2:B5")
        ws.merge_cells("A7:A9"); ws.merge_cells("B7:B9")
    p = _save(tmp_path, fill)
    ctx = register_workbook(p)
    plan = SheetRowPlanner().run(workbook_ctx=ctx, sheet="S")["plan"]
    anchors = [r for r in plan.rows if r.role is RowRole.ANCHOR]
    children = [r for r in plan.rows if r.role is RowRole.CHILD]
    totals = [r for r in plan.rows if r.role is RowRole.TOTAL]
    assert len(anchors) == 2
    assert len(children) == 5
    assert len(totals) == 2


def test_plan_row_per_pli_does_not_emit_kv_anchors(tmp_path):
    """KV anchors are reserved for SHEET_IS_PLI. Column-header labels like
    'STYLE' / 'COLOR' / 'ORDER QTY' must not leak into kv_anchors for a
    tabular sheet — otherwise the applier overwrites real column reads with
    adjacent-cell garbage (e.g., quantity := 'PLAN QTY')."""
    clear_cache()
    def fill(ws):
        ws["A1"] = "S NO"; ws["B1"] = "IO NO"; ws["F1"] = "STYLE"
        ws["K1"] = "COLOR"; ws["L1"] = "ORDER QTY"; ws["M1"] = "PLAN QTY"
        ws["A2"] = 1; ws["B2"] = 1063; ws["F2"] = "DWJE"
        ws["K2"] = "MAGENTA"; ws["L2"] = 2356; ws["M2"] = 2482
    p = _save(tmp_path, fill)
    ctx = register_workbook(p)
    plan = SheetRowPlanner().run(workbook_ctx=ctx, sheet="S")["plan"]
    assert plan.pli_mode is PliMode.ROW_PER_PLI
    assert plan.kv_anchors == []
