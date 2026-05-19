"""_collect_header_labels — populates header_labels for ROW_PER_PLI, skips claimed cols."""
import openpyxl
from app.enums.pli_mode import PliMode
from app.models.artifacts import HeaderLabel, SheetPlan, StageBandSpec, StageColumn
from app.services.planner.plan import _collect_header_labels


def _ws_with_header(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    # Row 3 header row
    ws["B3"] = "IO NO"
    ws["F3"] = "STYLE"
    ws["K3"] = "COLOR"
    ws["U3"] = "CUTTING"   # claimed by stage band
    ws["V3"] = "Actual"    # claimed sub-col
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    return openpyxl.load_workbook(p)["S"]


def test_collects_identity_columns_skipping_stage_cols(tmp_path) -> None:
    ws = _ws_with_header(tmp_path)
    band = StageBandSpec(
        name="b", name_cell="U3", sub_header_row=3,
        stage_columns=[StageColumn(name="CUTTING", name_cell="U3", primary_col="U",
                                    sub_columns={"Actual": "V"})],
    )
    plan = SheetPlan(sheet="S", pli_mode=PliMode.ROW_PER_PLI,
                     header_rows=[3], stage_bands=[band])
    labels = _collect_header_labels(ws, plan)
    raws = sorted(hl.raw for hl in labels)
    assert raws == ["COLOR", "IO NO", "STYLE"]
    assert all(isinstance(hl, HeaderLabel) for hl in labels)
    by_raw = {hl.raw: hl for hl in labels}
    assert by_raw["IO NO"].col == "B"
    assert by_raw["IO NO"].row == 3


def test_returns_empty_for_non_row_per_pli(tmp_path) -> None:
    ws = _ws_with_header(tmp_path)
    plan = SheetPlan(sheet="S", pli_mode=PliMode.SHEET_IS_PLI, header_rows=[3])
    assert _collect_header_labels(ws, plan) == []
    plan2 = SheetPlan(sheet="S", pli_mode=PliMode.SECTION_PER_PLI, header_rows=[3])
    assert _collect_header_labels(ws, plan2) == []
