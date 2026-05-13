from openpyxl import Workbook
from app.repositories.workbook_repo import register_workbook, clear_cache
from app.services.planner.surveyor import survey_sheet
from app.services.planner.row_classifier import classify_rows
from app.enums.row_role import RowRole


def _ctx_from(tmp_path, cells: dict[str, object], merges: list[str] | None = None):
    clear_cache()
    wb = Workbook(); ws = wb.active; ws.title = "S"
    for addr, val in cells.items():
        ws[addr] = val
    for mr in (merges or []):
        ws.merge_cells(mr)
    p = tmp_path / "x.xlsx"; wb.save(p)
    return register_workbook(p)


def test_classify_anchor_and_children(tmp_path):
    ctx = _ctx_from(
        tmp_path,
        {"A1": "IO NO", "B1": "COLOR",
         "A2": 1063, "B2": "MAGENTA",
         "B3": "NAVY", "B4": "WHITE"},
        merges=["A2:A4"],
    )
    sig = survey_sheet(ctx, "S")
    rows = classify_rows(ctx, "S", sig, identity_column="A")
    by_idx = {r.idx: r for r in rows}
    assert by_idx[1].role is RowRole.HEADER
    assert by_idx[2].role is RowRole.ANCHOR
    assert by_idx[3].role is RowRole.CHILD
    assert by_idx[3].anchor_idx == 2
    assert by_idx[4].role is RowRole.CHILD
    assert by_idx[4].anchor_idx == 2


def test_classify_total_via_sum_of_children(tmp_path):
    """A TOTAL row sums rows within the same group (anchor + its children)."""
    ctx = _ctx_from(
        tmp_path,
        {"A1": "IO NO", "C1": "QTY",
         "A2": 1063, "C2": 1000,
         "C3": 1500,         # CHILD of A2
         "C4": 2500},        # row 4 = sum of qty in group 0
        merges=["A2:A3"],
    )
    sig = survey_sheet(ctx, "S")
    rows = classify_rows(ctx, "S", sig, identity_column="A",
                        quantity_column_hint="C")
    by_idx = {r.idx: r for r in rows}
    assert by_idx[4].role is RowRole.TOTAL


def test_classify_grand_total_across_groups(tmp_path):
    """A GRAND_TOTAL row sums quantities across multiple groups."""
    ctx = _ctx_from(
        tmp_path,
        {"A1": "IO NO", "C1": "QTY",
         "A2": 1063, "C2": 1000,
         "A3": 1064, "C3": 2000,
         "C4": 3000},  # row 4 = grand sum across groups 0 and 1
    )
    sig = survey_sheet(ctx, "S")
    rows = classify_rows(ctx, "S", sig, identity_column="A",
                        quantity_column_hint="C")
    by_idx = {r.idx: r for r in rows}
    assert by_idx[4].role is RowRole.GRAND_TOTAL


def test_classify_blank_row(tmp_path):
    ctx = _ctx_from(
        tmp_path,
        {"A1": "IO NO", "A2": 1063, "A4": 1064},
    )
    sig = survey_sheet(ctx, "S")
    rows = classify_rows(ctx, "S", sig, identity_column="A")
    by_idx = {r.idx: r for r in rows}
    assert by_idx[3].role is RowRole.BLANK


def test_classify_repeat_header(tmp_path):
    ctx = _ctx_from(
        tmp_path,
        {"A1": "IO NO", "A2": 1063, "A3": "IO NO", "A4": 1064},
    )
    sig = survey_sheet(ctx, "S")
    rows = classify_rows(ctx, "S", sig, identity_column="A")
    by_idx = {r.idx: r for r in rows}
    assert by_idx[3].role is RowRole.REPEAT_HEADER
