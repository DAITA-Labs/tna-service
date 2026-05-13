from openpyxl import Workbook
from app.repositories.workbook_repo import register_workbook, clear_cache
from app.services.planner.surveyor import survey_sheet


def test_survey_basic_grid(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active
    ws.title = "S"
    ws["A1"] = "IO NO"; ws["B1"] = "STYLE"; ws["C1"] = "QTY"
    ws["A2"] = 1063;    ws["B2"] = "DWJE"; ws["C2"] = 2356
    ws["A3"] = 1064;    ws["B3"] = "DWJF"; ws["C3"] = 1500
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    sig = survey_sheet(ctx, "S")
    assert sig.sheet == "S"
    assert sig.max_row == 3
    assert "A" in sig.identity_col_candidates
    assert "IO NO" in sig.header_vocab_hits.get("A", [])


def test_survey_detects_merges(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active; ws.title = "S"
    ws["A1"] = "IO"; ws["A2"] = 1063
    ws.merge_cells("A2:A4")
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    sig = survey_sheet(ctx, "S")
    assert (2, 1, 4, 1) in sig.merges


def test_survey_detects_blank_run_gap(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active; ws.title = "S"
    ws["A1"] = "IO"
    ws["A2"] = 1063; ws["A3"] = 1064
    # rows 4, 5 entirely empty
    ws["A6"] = 1065
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    sig = survey_sheet(ctx, "S")
    assert (4, 5) in sig.blank_run_gaps


def test_survey_detects_kv_label_hits(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active; ws.title = "S"
    ws["A1"] = "Job No"; ws["B1"] = 63315
    ws["A2"] = "Quantity"; ws["B2"] = 254886
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    sig = survey_sheet(ctx, "S")
    labels = {x[0] for x in sig.kv_label_hits}
    assert "Job No" in labels and "Quantity" in labels
