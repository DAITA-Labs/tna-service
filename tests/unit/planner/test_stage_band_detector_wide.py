"""Stage-band detector — wide_sub_columns shape via stage_columns list."""
from datetime import date

from app.models.artifacts import SheetSignals
from app.repositories.workbook_repo import register_workbook
from app.services.planner.stage_band_detector import detect_stage_bands


def _signals(max_row: int, max_col: int) -> SheetSignals:
    return SheetSignals(sheet="S", max_row=max_row, max_col=max_col)


def test_wide_band_emits_stage_columns_list(tmp_path) -> None:
    """Detector populates stage_columns mirroring stage_cols (empty sub_columns)."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A3"] = "S NO"
    ws["B3"] = "IO"
    ws["U3"] = "CUTTING"
    ws["V3"] = "SEWING"
    ws["U4"] = date(2026, 3, 12)
    ws["V4"] = date(2026, 4, 3)
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    ctx = register_workbook(p)

    bands = detect_stage_bands(ctx, "S", _signals(max_row=4, max_col=22))
    assert len(bands) == 1
    b = bands[0]
    assert b.layout_mode == "wide_sub_columns"
    assert b.stage_cols == {"CUTTING": "U", "SEWING": "V"}
    assert len(b.stage_columns) == 2
    by_name = {sc.name: sc for sc in b.stage_columns}
    assert by_name["CUTTING"].primary_col == "U"
    assert by_name["CUTTING"].name_cell == "U3"
    assert by_name["CUTTING"].sub_columns == {}
    assert by_name["SEWING"].primary_col == "V"
