"""Tests for source_cell_verifier + header_match_verifier."""
from openpyxl import Workbook
from app.repositories.workbook_repo import register_workbook, clear_cache
from app.models.extraction import PLI, ExtractionResult
from app.services.validation.source_cell_verifier import SourceCellVerifier
from app.services.validation.header_match_verifier import HeaderMatchVerifier


def test_source_cell_verifier_passes_when_values_match(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active
    ws["K4"] = "131673"
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    pli = PLI(io_number="131673", source_sheet="Sheet",
              source_cells={"io_number": "K4"})
    result = ExtractionResult(plis=[pli], source_file=str(p))
    v = SourceCellVerifier(workbook_ctx=ctx)
    out = v.run(extraction=result)
    fs = out["findings"]
    assert not any(f.severity == "warn" for f in fs.findings)


def test_source_cell_verifier_warns_on_mismatch(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active
    ws["K4"] = "DIFFERENT"
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    pli = PLI(io_number="131673", source_sheet="Sheet",
              source_cells={"io_number": "K4"})
    result = ExtractionResult(plis=[pli], source_file=str(p))
    v = SourceCellVerifier(workbook_ctx=ctx)
    fs = v.run(extraction=result)["findings"]
    assert any(f.severity == "warn" and f.check == "source_cell" for f in fs.findings)


def test_header_match_verifier_passes(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active
    ws["K2"] = "Buyer Po No"
    ws["K4"] = "131673"
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    pli = PLI(io_number="131673", source_sheet="Sheet",
              source_cells={"io_number": "K4"})
    result = ExtractionResult(plis=[pli], source_file=str(p))
    v = HeaderMatchVerifier(workbook_ctx=ctx)
    fs = v.run(extraction=result)["findings"]
    assert not any(f.severity == "warn" and f.check == "header_match" for f in fs.findings)
