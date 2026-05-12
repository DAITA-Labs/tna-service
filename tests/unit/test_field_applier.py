"""Tests for app/services/applier/field_applier."""
from openpyxl import Workbook
from app.repositories.workbook_repo import register_workbook, clear_cache
from app.models.extraction import PLI
from app.models.artifacts import FieldMap, FieldLocation, PLIBoundaries
from app.enums.boundary_pattern import BoundaryPattern
from app.enums.location_pattern import LocationPattern
from app.services.applier.field_applier import (
    apply_field_map, is_real_pli, _PLI_IDENTITY_FIELDS,
)
import app.services.applier.patterns  # noqa: F401 — register handlers


def test_is_real_pli_requires_identity():
    assert is_real_pli(PLI(io_number="123"))
    assert is_real_pli(PLI(style_code="DWJE"))
    assert not is_real_pli(PLI(quantity=500))
    assert not is_real_pli(PLI())


def test_apply_field_map_records_source_cells(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active
    ws["K4"] = "131673"; ws["E4"] = "DWJE-1"
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    fm = FieldMap(sheet="Sheet", locations=[
        FieldLocation(field="io_number", pattern=LocationPattern.COLUMN, column="K",
                     data_start_row=4, data_end_row=4, confidence=1.0),
        FieldLocation(field="style_code", pattern=LocationPattern.COLUMN, column="E",
                     data_start_row=4, data_end_row=4, confidence=1.0),
    ])
    b = PLIBoundaries(sheet="Sheet", pattern=BoundaryPattern.ONE_ROW_PER_PLI,
                     data_start_row=4, data_end_row=4, confidence=1.0)
    plis = apply_field_map(ctx, "Sheet", fm, b)
    assert len(plis) == 1
    assert plis[0].source_cells["io_number"] == "K4"
    assert plis[0].source_cells["style_code"] == "E4"


def test_apply_field_map_strips_identity_on_repeat_header(tmp_path):
    """A row where canonical-field value equals a header label → strip id."""
    clear_cache()
    wb = Workbook(); ws = wb.active
    ws["B2"] = "IO NO"
    ws["B4"] = "1068"; ws["B5"] = "IO NO"; ws["B6"] = "1070"
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    fm = FieldMap(sheet="Sheet", locations=[
        FieldLocation(field="io_number", pattern=LocationPattern.COLUMN, column="B",
                     data_start_row=4, data_end_row=6, confidence=1.0),
    ])
    b = PLIBoundaries(sheet="Sheet", pattern=BoundaryPattern.ONE_ROW_PER_PLI,
                     data_start_row=4, data_end_row=6, confidence=1.0)
    plis = apply_field_map(ctx, "Sheet", fm, b)
    assert plis[0].io_number == "1068"
    assert plis[1].io_number is None
    assert plis[2].io_number == "1070"


def test_apply_field_map_vertical_merge_propagates(tmp_path):
    """In vertical_merge, identity columns propagate from the merge anchor."""
    clear_cache()
    wb = Workbook(); ws = wb.active
    ws["B4"] = "1063"
    ws.merge_cells("B4:B5")
    ws["K4"] = "MAGENTA"; ws["K5"] = "NAVY"
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    fm = FieldMap(sheet="Sheet", locations=[
        FieldLocation(field="io_number", pattern=LocationPattern.COLUMN, column="B",
                     data_start_row=4, data_end_row=5, confidence=1.0),
        FieldLocation(field="color_code", pattern=LocationPattern.COLUMN, column="K",
                     data_start_row=4, data_end_row=5, confidence=1.0),
    ])
    b = PLIBoundaries(sheet="Sheet", pattern=BoundaryPattern.VERTICAL_MERGE,
                     data_start_row=4, data_end_row=5,
                     grouping_columns=["B"], confidence=1.0)
    plis = apply_field_map(ctx, "Sheet", fm, b)
    assert plis[0].io_number == "1063"
    assert plis[1].io_number == "1063"
    assert plis[0].color_code == "MAGENTA"
    assert plis[1].color_code == "NAVY"
