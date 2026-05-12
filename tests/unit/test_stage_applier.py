"""Tests for app/services/applier/stage_applier."""
from datetime import datetime
from openpyxl import Workbook
from app.repositories.workbook_repo import register_workbook, clear_cache
from app.models.artifacts import (
    StageBandSet, StageBand, StageColumn, PLIBoundaries, FieldMap,
    FieldLocation, PLIMetadataLocation,
)
from app.enums.boundary_pattern import BoundaryPattern
from app.enums.location_pattern import LocationPattern
from app.enums.stage_layout_mode import StageLayoutMode
from app.services.applier.stage_applier import (
    apply_stage_band_set, strip_stage_columns_from_metadata,
)
import app.services.applier.patterns  # noqa: F401 — register handlers


def test_apply_wide_sub_columns(tmp_path):
    clear_cache()
    wb = Workbook(); ws = wb.active
    ws["R4"] = datetime(2026, 3, 25)
    ws["S4"] = datetime(2026, 3, 26)
    p = tmp_path / "x.xlsx"; wb.save(p)
    ctx = register_workbook(p)
    band = StageBand(
        section_name=None, name_row=2,
        layout_mode=StageLayoutMode.WIDE_SUB_COLUMNS,
        sub_header_row=3, data_start_row=4, data_end_row=4,
        stage_columns=[StageColumn(
            name="Trims Inhouse", name_cell="R2", primary_col="R",
            sub_columns={"actual": "S"},
        )],
        confidence=1.0,
    )
    sset = StageBandSet(sheet="Sheet", bands=[band], confidence=1.0)
    b = PLIBoundaries(sheet="Sheet", pattern=BoundaryPattern.ONE_ROW_PER_PLI,
                     data_start_row=4, data_end_row=4, confidence=1.0)
    out = apply_stage_band_set(ctx, "Sheet", sset, b)
    assert len(out) == 1
    assert out[0][0].name == "Trims Inhouse"
    assert out[0][0].planned_date == datetime(2026, 3, 25).date()
    assert out[0][0].metadata["actual"] == datetime(2026, 3, 26)


def test_strip_stage_columns_drops_overlap():
    fm = FieldMap(
        sheet="S",
        metadata_locations=[
            PLIMetadataLocation(key="cut_qty", pattern=LocationPattern.COLUMN,
                                column="Y", data_start_row=4, data_end_row=9,
                                confidence=0.9),
            PLIMetadataLocation(key="sewing_qty", pattern=LocationPattern.COLUMN,
                                column="AD", data_start_row=4, data_end_row=9,
                                confidence=0.9),
        ],
    )
    sset = StageBandSet(sheet="S", bands=[StageBand(
        section_name=None, name_row=2,
        layout_mode=StageLayoutMode.WIDE_SUB_COLUMNS,
        sub_header_row=3, data_start_row=4, data_end_row=9,
        stage_columns=[StageColumn(name="Sewing", name_cell="Z2",
                                  primary_col="Z", sub_columns={"qty": "AD"})],
        confidence=1.0,
    )], confidence=1.0)
    pruned = strip_stage_columns_from_metadata(fm, sset)
    keys = {m.key for m in pruned.metadata_locations}
    assert keys == {"cut_qty"}
