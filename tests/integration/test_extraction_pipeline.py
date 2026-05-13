"""End-to-end smoke test of the rewritten extraction pipeline with a fake LLM."""
from openpyxl import Workbook
from app.repositories.workbook_repo import clear_cache
from app.services.extraction import extract
from app.models.artifacts import (
    CanonicalNameMap, LayoutHints, PlanVerdict,
)
from app.services.agents.sheet_classifier import SheetClassifierOutput


class _FakeLLM:
    """Stub that satisfies LLMProvider Protocol — returns canned Pydantic outputs."""

    model: str = "fake-llm"

    def complete_with_schema(self, *, system, user, output_schema, tool_name):
        name = output_schema.__name__
        if name == "SheetClassifierOutput":
            return SheetClassifierOutput(relevant_sheets=["S"])
        if name == "CanonicalNameMap":
            return CanonicalNameMap(
                field_labels={"IO NO": "io_number", "COLOR": "color_code"},
                stage_names={},
            )
        if name == "LayoutHints":
            return LayoutHints()
        if name == "PlanVerdict":
            return PlanVerdict(verdict="looks_correct")
        raise NotImplementedError(f"_FakeLLM: unknown schema {name!r}")


def test_extract_christian_berg_like(tmp_path):
    clear_cache()
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "IO NO"
    ws["B1"] = "COLOR"
    ws["A2"] = "1063"
    ws["B2"] = "MAGENTA"
    ws["B3"] = "NAVY"
    ws.merge_cells("A2:A3")
    p = tmp_path / "x.xlsx"
    wb.save(p)

    result = extract(p, llm=_FakeLLM())
    assert len(result.plis) == 2
    assert result.plis[0].io_number == "1063"
    assert result.plis[1].io_number == "1063"
